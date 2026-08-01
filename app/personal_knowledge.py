from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Iterable

from pydantic import BaseModel, Field, ValidationError, field_validator

from .audit import emit
from .config import get_settings
from .database import db_session
from .processes import MANUFACTURING_STEPS, PROCESS_ALIASES, process_query_terms


PARSER_VERSION = "personal-v1"
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".txt", ".csv", ".xlsx"}
MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Canonical, reviewable vocabulary used for deterministic metadata extraction.
DOCUMENT_TYPE_TAGS = {
    "SOP": ("sop", "standard operating procedure", "标准作业程序", "作业指导书"),
    "CONTROL PLAN": ("control plan", "控制计划"),
    "GUIDE": ("guide", "guideline", "指南", "指引"),
    "REPORT": ("report", "报告"),
    "RECORD": ("record", "记录"),
    "TRAINING": ("training", "培训"),
}
TOPIC_TAGS = {
    "SCRATCH": ("scratch", "划伤", "刮伤"),
    "PARTICLE": ("particle", "颗粒"),
    "MAINTENANCE": ("maintenance", "维护", "保养"),
    "SPC": ("spc", "statistical process control", "统计过程控制"),
    "YIELD": ("yield", "良率"),
    "CPK": ("cpk", "process capability index", "过程能力指数"),
    "RECIPE": ("recipe", "配方"),
    "ALARM": ("alarm", "报警", "告警"),
}
MAX_TAG_SOURCE_CHARS = 200_000


class PersonalKnowledgeError(Exception):
    def __init__(self, status_code: int, error_code: str, detail: str, retryable: bool = False):
        super().__init__(detail)
        self.status_code = status_code
        self.error_code = error_code
        self.detail = detail
        self.retryable = retryable


class PersonalDocumentMetadata(BaseModel):
    title: str | None = Field(default=None, max_length=200)
    process_steps: list[str] = Field(default_factory=list, max_length=20)
    equipment_ids: list[str] = Field(default_factory=list, max_length=100)
    effective_date: date | None = None
    tags: list[str] = Field(default_factory=list, max_length=30)

    @field_validator("title")
    @classmethod
    def clean_title(cls, value: str | None) -> str | None:
        cleaned = " ".join((value or "").strip().split())
        return cleaned or None

    @field_validator("process_steps")
    @classmethod
    def clean_process_steps(cls, values: list[str]) -> list[str]:
        cleaned = list(dict.fromkeys(value.strip().upper() for value in values if value.strip()))
        invalid = [value for value in cleaned if value not in MANUFACTURING_STEPS]
        if invalid:
            raise ValueError(f"不支持的制程: {', '.join(invalid)}")
        return cleaned

    @field_validator("equipment_ids")
    @classmethod
    def clean_equipment_ids(cls, values: list[str]) -> list[str]:
        cleaned = list(dict.fromkeys(value.strip().upper() for value in values if value.strip()))
        if any(not re.fullmatch(r"[A-Z][A-Z0-9_-]{1,31}", value) for value in cleaned):
            raise ValueError("设备编号格式不正确")
        return cleaned

    @field_validator("tags")
    @classmethod
    def clean_tags(cls, values: list[str]) -> list[str]:
        cleaned: list[str] = []
        seen: set[str] = set()
        for value in values:
            tag = " ".join(value.strip().split())
            key = tag.casefold()
            if tag and key not in seen:
                cleaned.append(tag)
                seen.add(key)
        if any(len(value) > 40 for value in cleaned):
            raise ValueError("单个标签不能超过 40 个字符")
        return cleaned


class PersonalDocumentPatch(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=200)
    process_steps: list[str] | None = Field(default=None, max_length=20)
    equipment_ids: list[str] | None = Field(default=None, max_length=100)
    effective_date: date | None = None
    tags: list[str] | None = Field(default=None, max_length=30)

    @field_validator("title")
    @classmethod
    def clean_title(cls, value: str | None) -> str | None:
        if value is None:
            return None
        cleaned = " ".join(value.strip().split())
        if not cleaned:
            raise ValueError("标题不能为空")
        return cleaned

    @field_validator("process_steps")
    @classmethod
    def clean_process_steps(cls, values: list[str] | None) -> list[str] | None:
        return None if values is None else PersonalDocumentMetadata.clean_process_steps(values)

    @field_validator("equipment_ids")
    @classmethod
    def clean_equipment_ids(cls, values: list[str] | None) -> list[str] | None:
        return None if values is None else PersonalDocumentMetadata.clean_equipment_ids(values)

    @field_validator("tags")
    @classmethod
    def clean_tags(cls, values: list[str] | None) -> list[str] | None:
        return None if values is None else PersonalDocumentMetadata.clean_tags(values)


@dataclass(frozen=True)
class ExtractedSection:
    page_number: int | None
    section_path: str
    content: str


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def _json_list(value: str | None) -> list[str]:
    try:
        result = json.loads(value or "[]")
    except json.JSONDecodeError:
        return []
    return result if isinstance(result, list) else []


def _document_dict(row) -> dict:
    item = dict(row)
    item["process_steps"] = _json_list(item.pop("process_steps_json", "[]"))
    item["equipment_ids"] = _json_list(item.pop("equipment_ids_json", "[]"))
    item["tags"] = _json_list(item.pop("tags_json", "[]"))
    item["tags_source"] = item.get("tags_source", "auto")
    item.pop("owner_user_id", None)
    item.pop("storage_key", None)
    item.pop("idempotency_key", None)
    return item


def _normalise_tags(values: Iterable[str]) -> list[str]:
    """Return stable, case-insensitively unique tags within the public limits."""
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        tag = " ".join(str(value).strip().split())
        key = tag.casefold()
        if tag and len(tag) <= 40 and key not in seen:
            result.append(tag)
            seen.add(key)
        if len(result) == 30:
            break
    return result


def _automatic_tags(title: str, filename: str, sections: Iterable[ExtractedSection]) -> list[str]:
    """Derive explainable tags from document text without a model call."""
    parts = [title, filename]
    remaining = MAX_TAG_SOURCE_CHARS - sum(len(part) for part in parts)
    for section in sections:
        if remaining <= 0:
            break
        part = f"{section.section_path}\n{section.content}"
        parts.append(part[:remaining])
        remaining -= len(part)
    text = "\n".join(parts)
    compact = text.casefold().replace(" ", "")
    tags: list[str] = []
    for step in MANUFACTURING_STEPS:
        aliases = PROCESS_ALIASES[step]
        if step.casefold() in text.casefold() or any(alias.casefold().replace(" ", "") in compact for alias in aliases):
            tags.append(step)
    # Equipment identifiers intentionally follow the same conservative contract as metadata,
    # while requiring a digit, dash or underscore so ordinary prose is not labelled equipment.
    tags.extend(re.findall(r"\b(?=[A-Za-z][A-Za-z0-9_-]{1,31}\b)(?=[A-Za-z0-9_-]*[0-9_-])[A-Za-z][A-Za-z0-9_-]+\b", text))
    suffix = Path(filename).suffix.lstrip(".").upper()
    for tag, aliases in DOCUMENT_TYPE_TAGS.items():
        if any(alias.casefold().replace(" ", "") in compact for alias in aliases):
            tags.append(tag)
    if suffix:
        tags.append(suffix)
    for tag, aliases in TOPIC_TAGS.items():
        if any(alias.casefold().replace(" ", "") in compact for alias in aliases):
            tags.append(tag)
    return _normalise_tags(tags)


def parse_metadata(raw: str | None) -> PersonalDocumentMetadata:
    try:
        return PersonalDocumentMetadata.model_validate_json(raw or "{}")
    except ValidationError as exc:
        fields = [".".join(str(part) for part in error["loc"]) for error in exc.errors()]
        raise PersonalKnowledgeError(400, "INVALID_METADATA", f"元数据不符合要求: {', '.join(fields)}") from exc


def _safe_filename(filename: str | None) -> str:
    cleaned = (filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not cleaned or len(cleaned) > 255:
        raise PersonalKnowledgeError(422, "INVALID_FILE_CONTENT", "文件名为空或过长")
    return cleaned


def _validate_file(filename: str, data: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise PersonalKnowledgeError(415, "UNSUPPORTED_FILE_TYPE", "当前不支持该文件类型")
    if not data:
        raise PersonalKnowledgeError(422, "INVALID_FILE_CONTENT", "文件内容为空")
    try:
        if suffix == ".pdf" and not data.startswith(b"%PDF-"):
            raise ValueError("invalid PDF signature")
        if suffix in {".docx", ".xlsx"}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                expanded_size = sum(item.file_size for item in archive.infolist())
                compressed_size = max(1, sum(item.compress_size for item in archive.infolist()))
            expected = "word/document.xml" if suffix == ".docx" else "xl/workbook.xml"
            if expected not in names or expanded_size > 200 * 1024 * 1024 or expanded_size / compressed_size > 100:
                raise ValueError("invalid Office document signature")
        if suffix in {".md", ".txt", ".csv"}:
            if b"\x00" in data[:8192]:
                raise ValueError("binary content")
            data.decode("utf-8-sig")
    except (UnicodeDecodeError, ValueError, zipfile.BadZipFile) as exc:
        raise PersonalKnowledgeError(422, "INVALID_FILE_CONTENT", "文件内容与声明类型不匹配或已经损坏") from exc
    return suffix, MIME_TYPES[suffix]


def _storage_path(storage_key: str) -> Path:
    root = get_settings().personal_knowledge_storage_path.resolve()
    target = (root / storage_key).resolve()
    if root != target and root not in target.parents:
        raise PersonalKnowledgeError(500, "STORAGE_PATH_ERROR", "文件存储路径不安全")
    return target


def _latest_job(connection, document_id: str, owner_user_id: str):
    return connection.execute(
        """SELECT job_id, document_id, status, stage, progress, attempt,
                  error_code, error_message, retryable, created_at, updated_at
           FROM personal_knowledge_jobs
           WHERE document_id = ? AND owner_user_id = ?
           ORDER BY attempt DESC, created_at DESC, job_id DESC LIMIT 1""",
        (document_id, owner_user_id),
    ).fetchone()


def _initial_upload_job(connection, document_id: str, owner_user_id: str):
    """Return the immutable first upload attempt for an idempotent replay."""
    return connection.execute(
        """SELECT job_id, document_id, status, stage, progress, attempt,
                  error_code, error_message, retryable, created_at, updated_at
           FROM personal_knowledge_jobs
           WHERE document_id = ? AND owner_user_id = ?
           ORDER BY attempt ASC, created_at ASC, job_id ASC LIMIT 1""",
        (document_id, owner_user_id),
    ).fetchone()


def _job_dict(row) -> dict:
    item = dict(row)
    error_code = item.pop("error_code", None)
    error_message = item.pop("error_message", None)
    retryable = bool(item.pop("retryable", 0))
    item["error"] = (
        {"error_code": error_code, "detail": error_message, "retryable": retryable}
        if error_code
        else None
    )
    return item


def create_personal_upload(
    owner_user_id: str,
    filename: str | None,
    data: bytes,
    metadata_raw: str | None,
    idempotency_key: str,
) -> dict:
    settings = get_settings()
    if len(data) > settings.personal_knowledge_max_upload_bytes:
        raise PersonalKnowledgeError(413, "FILE_TOO_LARGE", "文件超过当前上传大小限制")
    safe_name = _safe_filename(filename)
    suffix, mime_type = _validate_file(safe_name, data)
    metadata = parse_metadata(metadata_raw)
    tags_source = "user" if "tags" in metadata.model_fields_set else "auto"
    checksum = hashlib.sha256(data).hexdigest()

    with db_session() as connection:
        idempotent = connection.execute(
            """SELECT * FROM personal_knowledge_documents
               WHERE owner_user_id = ? AND idempotency_key = ?""",
            (owner_user_id, idempotency_key),
        ).fetchone()
        if idempotent is not None:
            if idempotent["checksum_sha256"] != checksum:
                raise PersonalKnowledgeError(409, "IDEMPOTENCY_CONFLICT", "同一幂等键不能用于不同文件")
            # A replay identifies the original upload transaction, not a later
            # retry.  Document detail continues to expose the latest job.
            job = _initial_upload_job(connection, idempotent["document_id"], owner_user_id)
            response = _upload_response(idempotent, job)
            response["_start_processing"] = False
            return response
        duplicate = connection.execute(
            """SELECT document_id FROM personal_knowledge_documents
               WHERE owner_user_id = ? AND checksum_sha256 = ?""",
            (owner_user_id, checksum),
        ).fetchone()
        if duplicate is not None:
            raise PersonalKnowledgeError(409, "DUPLICATE_DOCUMENT", "当前个人知识库已存在完全相同的文件")

    document_id = str(uuid.uuid4())
    job_id = str(uuid.uuid4())
    owner_directory = hashlib.sha256(owner_user_id.encode("utf-8")).hexdigest()[:24]
    storage_key = f"{owner_directory}/{document_id}/original{suffix}"
    target = _storage_path(storage_key)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + ".part")
    temporary.write_bytes(data)
    temporary.replace(target)
    timestamp = _now()
    title = metadata.title or Path(safe_name).stem[:200] or "未命名文档"
    try:
        with db_session() as connection:
            connection.execute(
                """INSERT INTO knowledge.personal_knowledge_documents(
                       document_id, owner_user_id, filename, title, mime_type, size_bytes,
                       checksum_sha256, process_steps_json, equipment_ids_json, effective_date,
                       tags_json, tags_source, status, storage_key, parser_version, idempotency_key,
                       created_at, updated_at
                   ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'uploaded', ?, ?, ?, ?, ?)""",
                (
                    document_id, owner_user_id, safe_name, title, mime_type, len(data), checksum,
                    json.dumps(metadata.process_steps), json.dumps(metadata.equipment_ids),
                    metadata.effective_date.isoformat() if metadata.effective_date else None,
                    json.dumps(_normalise_tags(metadata.tags), ensure_ascii=False), tags_source, storage_key, PARSER_VERSION,
                    idempotency_key, timestamp, timestamp,
                ),
            )
            connection.execute(
                """INSERT INTO knowledge.personal_knowledge_jobs(
                       job_id, document_id, owner_user_id, status, stage, progress, attempt,
                       created_at, updated_at
                   ) VALUES (?, ?, ?, 'uploaded', 'queued', 0, 1, ?, ?)""",
                (job_id, document_id, owner_user_id, timestamp, timestamp),
            )
            connection.commit()
    except Exception:
        target.unlink(missing_ok=True)
        raise

    emit(
        "personal_knowledge_uploaded",
        job_id,
        status="success",
        document_id=document_id,
        owner_fingerprint=hashlib.sha256(owner_user_id.encode("utf-8")).hexdigest()[:12],
        filename_suffix=suffix,
        size_bytes=len(data),
    )
    with db_session() as connection:
        document = connection.execute(
            "SELECT * FROM personal_knowledge_documents WHERE document_id = ?", (document_id,)
        ).fetchone()
        job = _latest_job(connection, document_id, owner_user_id)
    response = _upload_response(document, job)
    response["_start_processing"] = True
    return response


def _upload_response(document, job) -> dict:
    return {
        "document_id": document["document_id"],
        "job_id": job["job_id"],
        "status": document["status"],
        "filename": document["filename"],
        "status_url": f"/api/knowledge/jobs/{job['job_id']}",
    }


def _update_progress(
    document_id: str,
    owner_user_id: str,
    job_id: str,
    status: str,
    stage: str,
    progress: int,
) -> None:
    timestamp = _now()
    with db_session() as connection:
        connection.execute(
            """UPDATE knowledge.personal_knowledge_documents
               SET status = ?, error_code = NULL, error_message = NULL, updated_at = ?
               WHERE document_id = ? AND owner_user_id = ?""",
            (status, timestamp, document_id, owner_user_id),
        )
        connection.execute(
            """UPDATE knowledge.personal_knowledge_jobs
               SET status = ?, stage = ?, progress = ?, error_code = NULL,
                   error_message = NULL, retryable = 0, updated_at = ?
               WHERE job_id = ? AND owner_user_id = ?""",
            (status, stage, progress, timestamp, job_id, owner_user_id),
        )
        connection.commit()


def _mark_failed(
    document_id: str,
    owner_user_id: str,
    job_id: str,
    error_code: str,
    detail: str,
    retryable: bool,
) -> None:
    timestamp = _now()
    with db_session() as connection:
        connection.execute(
            """UPDATE knowledge.personal_knowledge_documents
               SET status = 'failed', error_code = ?, error_message = ?, updated_at = ?
               WHERE document_id = ? AND owner_user_id = ?""",
            (error_code, detail, timestamp, document_id, owner_user_id),
        )
        connection.execute(
            """UPDATE knowledge.personal_knowledge_jobs
               SET status = 'failed', stage = 'failed', error_code = ?, error_message = ?,
                   retryable = ?, updated_at = ?
               WHERE job_id = ? AND owner_user_id = ?""",
            (error_code, detail, int(retryable), timestamp, job_id, owner_user_id),
        )
        connection.commit()


def _text_sections(text: str, default_title: str, markdown: bool = False) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    heading = default_title
    body: list[str] = []
    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.rstrip()
        is_heading = markdown and bool(re.match(r"^#{1,6}\s+\S", line))
        if is_heading:
            content = "\n".join(body).strip()
            if content:
                sections.append(ExtractedSection(None, heading, content))
            heading = re.sub(r"^#{1,6}\s+", "", line).strip()
            body = []
        else:
            body.append(line)
    content = "\n".join(body).strip()
    if content:
        sections.append(ExtractedSection(None, heading, content))
    return sections


def _extract_sections(path: Path, suffix: str, title: str) -> list[ExtractedSection]:
    if suffix in {".txt", ".md"}:
        return _text_sections(path.read_text(encoding="utf-8-sig"), title, suffix == ".md")
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = ["\t".join(cell.strip() for cell in row) for row in csv.reader(handle)]
        return [ExtractedSection(None, title, "\n".join(row for row in rows if row.strip()).strip())]
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise PersonalKnowledgeError(422, "INVALID_FILE_CONTENT", "PDF 已加密，无法解析")
        sections = [
            ExtractedSection(index, f"第 {index} 页", (page.extract_text() or "").strip())
            for index, page in enumerate(reader.pages, start=1)
        ]
        if sum(len(section.content) for section in sections) < 20:
            raise PersonalKnowledgeError(422, "OCR_REQUIRED", "该文件需要 OCR，当前版本暂不支持")
        return [section for section in sections if section.content]
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        sections: list[ExtractedSection] = []
        heading = title
        body: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                if body:
                    sections.append(ExtractedSection(None, heading, "\n".join(body)))
                heading, body = text, []
            else:
                body.append(text)
        if body:
            sections.append(ExtractedSection(None, heading, "\n".join(body)))
        for index, table in enumerate(document.tables, start=1):
            rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            content = "\n".join(row for row in rows if row.strip()).strip()
            if content:
                sections.append(ExtractedSection(None, f"表格 {index}", content))
        return sections
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sections = []
        try:
            for worksheet in workbook.worksheets:
                rows = []
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(value.strip() for value in values):
                        rows.append("\t".join(values).rstrip())
                if rows:
                    sections.append(ExtractedSection(None, worksheet.title, "\n".join(rows)))
        finally:
            workbook.close()
        return sections
    raise PersonalKnowledgeError(415, "UNSUPPORTED_FILE_TYPE", "当前不支持该文件类型")


def _split_content(content: str, max_chars: int = 2600, overlap_chars: int = 260) -> list[str]:
    normalized = re.sub(r"\n{3,}", "\n\n", content).strip()
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(len(normalized), start + max_chars)
        if end < len(normalized):
            boundary = max(normalized.rfind("\n", start + max_chars // 2, end), normalized.rfind("。", start + max_chars // 2, end))
            if boundary > start:
                end = boundary + 1
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(start + 1, end - overlap_chars)
    return chunks


def process_personal_document(document_id: str, owner_user_id: str, job_id: str) -> None:
    try:
        _update_progress(document_id, owner_user_id, job_id, "parsing", "extracting_text", 10)
        with db_session() as connection:
            document = connection.execute(
                """SELECT * FROM personal_knowledge_documents
                   WHERE document_id = ? AND owner_user_id = ?""",
                (document_id, owner_user_id),
            ).fetchone()
        if document is None:
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
        path = _storage_path(document["storage_key"])
        if not path.exists():
            raise PersonalKnowledgeError(500, "STORAGE_FILE_MISSING", "原始文件不存在", True)
        sections = _extract_sections(path, path.suffix.lower(), document["title"])
        automatic_tags = _automatic_tags(document["title"], document["filename"], sections)
        _update_progress(document_id, owner_user_id, job_id, "chunking", "chunking", 45)
        chunks: list[tuple[int | None, str, str]] = []
        for section in sections:
            chunks.extend(
                (section.page_number, section.section_path, content)
                for content in _split_content(section.content)
            )
        if not chunks:
            raise PersonalKnowledgeError(422, "INVALID_FILE_CONTENT", "文件中没有可索引文本")
        _update_progress(document_id, owner_user_id, job_id, "indexing", "full_text_index", 75)
        timestamp = _now()
        with db_session() as connection:
            connection.execute(
                "DELETE FROM knowledge.personal_knowledge_fts WHERE document_id = ?", (document_id,)
            )
            connection.execute(
                "DELETE FROM knowledge.personal_knowledge_chunks WHERE document_id = ?", (document_id,)
            )
            for index, (page_number, section_path, content) in enumerate(chunks):
                chunk_id = str(uuid.uuid4())
                cursor = connection.execute(
                    """INSERT INTO knowledge.personal_knowledge_chunks(
                           chunk_id, document_id, owner_user_id, page_number, section_path,
                           chunk_index, content, token_count, created_at
                       ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        chunk_id, document_id, owner_user_id, page_number, section_path,
                        index, content, max(1, len(content) // 2), timestamp,
                    ),
                )
                connection.execute(
                    """INSERT INTO knowledge.personal_knowledge_fts(
                           rowid, title, section_path, content, owner_user_id, document_id, chunk_id
                       ) VALUES (?, ?, ?, ?, ?, ?, ?)""",
                    (cursor.lastrowid, document["title"], section_path, content, owner_user_id, document_id, chunk_id),
                )
            connection.execute(
                """UPDATE knowledge.personal_knowledge_documents
                   SET status = 'ready', error_code = NULL, error_message = NULL,
                       tags_json = CASE WHEN tags_source = 'auto' THEN ? ELSE tags_json END,
                       updated_at = ?
                   WHERE document_id = ? AND owner_user_id = ?""",
                (json.dumps(automatic_tags, ensure_ascii=False), timestamp, document_id, owner_user_id),
            )
            connection.execute(
                """UPDATE knowledge.personal_knowledge_jobs
                   SET status = 'ready', stage = 'ready', progress = 100, error_code = NULL,
                       error_message = NULL, retryable = 0, updated_at = ?
                   WHERE job_id = ? AND owner_user_id = ?""",
                (timestamp, job_id, owner_user_id),
            )
            connection.commit()
        emit(
            "personal_knowledge_ingestion_finished",
            job_id,
            status="success",
            document_id=document_id,
            chunks=len(chunks),
        )
    except PersonalKnowledgeError as exc:
        _mark_failed(document_id, owner_user_id, job_id, exc.error_code, exc.detail, exc.retryable)
        emit("personal_knowledge_ingestion_finished", job_id, status="error", document_id=document_id, error_code=exc.error_code)
    except Exception as exc:
        _mark_failed(document_id, owner_user_id, job_id, "INGESTION_FAILED", "文档解析或索引失败", True)
        emit(
            "personal_knowledge_ingestion_finished",
            job_id,
            status="error",
            document_id=document_id,
            error_code="INGESTION_FAILED",
            error_type=type(exc).__name__,
        )


def get_personal_job(job_id: str, owner_user_id: str) -> dict:
    with db_session() as connection:
        row = connection.execute(
            """SELECT job_id, document_id, status, stage, progress, attempt,
                      error_code, error_message, retryable, created_at, updated_at
               FROM personal_knowledge_jobs WHERE job_id = ? AND owner_user_id = ?""",
            (job_id, owner_user_id),
        ).fetchone()
    if row is None:
        raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "任务不存在")
    return _job_dict(row)


def _encode_cursor(updated_at: str, document_id: str) -> str:
    raw = json.dumps([updated_at, document_id], separators=(",", ":")).encode("utf-8")
    return base64.urlsafe_b64encode(raw).decode("ascii").rstrip("=")


def _decode_cursor(cursor: str) -> tuple[str, str]:
    try:
        padded = cursor + "=" * (-len(cursor) % 4)
        value = json.loads(base64.urlsafe_b64decode(padded).decode("utf-8"))
        if not isinstance(value, list) or len(value) != 2 or not all(isinstance(item, str) for item in value):
            raise ValueError
        return value[0], value[1]
    except (ValueError, json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise PersonalKnowledgeError(400, "INVALID_CURSOR", "分页游标无效") from exc


def list_personal_documents(
    owner_user_id: str,
    query: str = "",
    status: str | None = None,
    limit: int = 20,
    cursor: str | None = None,
) -> dict:
    clauses = ["owner_user_id = ?"]
    params: list[object] = [owner_user_id]
    if status:
        clauses.append("status = ?")
        params.append(status)
    if query.strip():
        clauses.append("""(title LIKE ? OR filename LIKE ?
            OR EXISTS (SELECT 1 FROM json_each(tags_json) value WHERE value.value LIKE ?)
            OR EXISTS (SELECT 1 FROM json_each(process_steps_json) value WHERE value.value LIKE ?)
            OR EXISTS (SELECT 1 FROM json_each(equipment_ids_json) value WHERE value.value LIKE ?))""")
        value = f"%{query.strip()}%"
        params.extend([value, value, value, value, value])
    if cursor:
        cursor_time, cursor_id = _decode_cursor(cursor)
        clauses.append("(updated_at < ? OR (updated_at = ? AND document_id < ?))")
        params.extend([cursor_time, cursor_time, cursor_id])
    bounded = max(1, min(limit, 100))
    params.append(bounded + 1)
    with db_session() as connection:
        rows = connection.execute(
            f"""SELECT d.*, (SELECT COUNT(*) FROM personal_knowledge_chunks c
                              WHERE c.document_id = d.document_id) AS chunk_count
                 FROM personal_knowledge_documents d
                 WHERE {' AND '.join(clauses)}
                 ORDER BY updated_at DESC, document_id DESC LIMIT ?""",
            params,
        ).fetchall()
    has_more = len(rows) > bounded
    selected = rows[:bounded]
    return {
        "items": [_document_dict(row) for row in selected],
        "next_cursor": _encode_cursor(selected[-1]["updated_at"], selected[-1]["document_id"]) if has_more else None,
    }


def get_personal_document(document_id: str, owner_user_id: str, include_chunks: bool = False) -> dict:
    with db_session() as connection:
        row = connection.execute(
            """SELECT d.*, (SELECT COUNT(*) FROM personal_knowledge_chunks c
                              WHERE c.document_id = d.document_id) AS chunk_count
               FROM personal_knowledge_documents d
               WHERE d.document_id = ? AND d.owner_user_id = ?""",
            (document_id, owner_user_id),
        ).fetchone()
        if row is None:
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
        result = _document_dict(row)
        if include_chunks and row["status"] == "deleting":
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
        job = _latest_job(connection, document_id, owner_user_id)
        result["latest_job"] = _job_dict(job) if job else None
        if include_chunks:
            chunks = connection.execute(
                """SELECT chunk_id, page_number, section_path, chunk_index, content, token_count
                   FROM personal_knowledge_chunks
                   WHERE document_id = ? AND owner_user_id = ? ORDER BY chunk_index""",
                (document_id, owner_user_id),
            ).fetchall()
            result["chunks"] = [dict(chunk) for chunk in chunks]
    return result


def update_personal_document(document_id: str, owner_user_id: str, patch: PersonalDocumentPatch) -> dict:
    updates: dict[str, object] = {}
    fields = patch.model_fields_set
    if "title" in fields:
        updates["title"] = patch.title
    if "process_steps" in fields:
        updates["process_steps_json"] = json.dumps(patch.process_steps or [])
    if "equipment_ids" in fields:
        updates["equipment_ids_json"] = json.dumps(patch.equipment_ids or [])
    if "effective_date" in fields:
        updates["effective_date"] = patch.effective_date.isoformat() if patch.effective_date else None
    if "tags" in fields:
        updates["tags_json"] = json.dumps(_normalise_tags(patch.tags or []), ensure_ascii=False)
        updates["tags_source"] = "user"
    if not updates:
        return get_personal_document(document_id, owner_user_id)
    with db_session() as connection:
        existing = connection.execute(
            """SELECT status FROM personal_knowledge_documents
               WHERE document_id = ? AND owner_user_id = ?""",
            (document_id, owner_user_id),
        ).fetchone()
        if existing is None:
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
        if existing["status"] == "deleting":
            raise PersonalKnowledgeError(409, "INVALID_DOCUMENT_STATE", "删除中的文档不能修改")
        updates["updated_at"] = _now()
        assignments = ", ".join(f"{column} = ?" for column in updates)
        connection.execute(
            f"""UPDATE knowledge.personal_knowledge_documents SET {assignments}
                 WHERE document_id = ? AND owner_user_id = ?""",
            [*updates.values(), document_id, owner_user_id],
        )
        if "title" in updates:
            connection.execute(
                """UPDATE knowledge.personal_knowledge_fts SET title = ?
                   WHERE document_id = ? AND owner_user_id = ?""",
                (updates["title"], document_id, owner_user_id),
            )
        connection.commit()
    return get_personal_document(document_id, owner_user_id)


def retry_personal_document(document_id: str, owner_user_id: str) -> dict:
    timestamp = _now()
    with db_session() as connection:
        document = connection.execute(
            """SELECT * FROM personal_knowledge_documents
               WHERE document_id = ? AND owner_user_id = ?""",
            (document_id, owner_user_id),
        ).fetchone()
        if document is None:
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
        if document["status"] != "failed":
            raise PersonalKnowledgeError(409, "INVALID_DOCUMENT_STATE", "只有失败的文档可以重试")
        if document["error_code"] == "DELETE_FAILED":
            raise PersonalKnowledgeError(409, "INVALID_DOCUMENT_STATE", "删除失败的文档应重新执行删除")
        attempt = connection.execute(
            "SELECT COALESCE(MAX(attempt), 0) + 1 FROM personal_knowledge_jobs WHERE document_id = ?",
            (document_id,),
        ).fetchone()[0]
        job_id = str(uuid.uuid4())
        connection.execute(
            """UPDATE knowledge.personal_knowledge_documents
               SET status = 'uploaded', error_code = NULL, error_message = NULL, updated_at = ?
               WHERE document_id = ? AND owner_user_id = ?""",
            (timestamp, document_id, owner_user_id),
        )
        connection.execute(
            """INSERT INTO knowledge.personal_knowledge_jobs(
                   job_id, document_id, owner_user_id, status, stage, progress, attempt,
                   created_at, updated_at
               ) VALUES (?, ?, ?, 'uploaded', 'queued', 0, ?, ?, ?)""",
            (job_id, document_id, owner_user_id, attempt, timestamp, timestamp),
        )
        connection.commit()
    return {"document_id": document_id, "job_id": job_id, "status": "uploaded", "status_url": f"/api/knowledge/jobs/{job_id}"}


def begin_delete_personal_document(document_id: str, owner_user_id: str) -> dict:
    timestamp = _now()
    job_id = str(uuid.uuid4())
    with db_session() as connection:
        document = connection.execute(
            """SELECT status FROM personal_knowledge_documents
               WHERE document_id = ? AND owner_user_id = ?""",
            (document_id, owner_user_id),
        ).fetchone()
        if document is None:
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
        if document["status"] == "deleting":
            raise PersonalKnowledgeError(409, "INVALID_DOCUMENT_STATE", "文档已经在删除中")
        connection.execute(
            """UPDATE knowledge.personal_knowledge_documents SET status = 'deleting', updated_at = ?
               WHERE document_id = ? AND owner_user_id = ?""",
            (timestamp, document_id, owner_user_id),
        )
        connection.execute(
            """INSERT INTO knowledge.personal_knowledge_jobs(
                   job_id, document_id, owner_user_id, status, stage, progress, attempt,
                   created_at, updated_at
               ) VALUES (?, ?, ?, 'deleting', 'deleting', 0, 1, ?, ?)""",
            (job_id, document_id, owner_user_id, timestamp, timestamp),
        )
        connection.commit()
    return {"document_id": document_id, "job_id": job_id, "status": "deleting"}


def process_personal_document_deletion(document_id: str, owner_user_id: str, job_id: str) -> None:
    try:
        with db_session() as connection:
            document = connection.execute(
                """SELECT storage_key FROM personal_knowledge_documents
                   WHERE document_id = ? AND owner_user_id = ?""",
                (document_id, owner_user_id),
            ).fetchone()
        if document is None:
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
        path = _storage_path(document["storage_key"])
        path.unlink(missing_ok=True)
        for directory in (path.parent, path.parent.parent):
            try:
                directory.rmdir()
            except OSError:
                pass
        timestamp = _now()
        with db_session() as connection:
            connection.execute(
                "DELETE FROM knowledge.personal_knowledge_fts WHERE document_id = ? AND owner_user_id = ?",
                (document_id, owner_user_id),
            )
            connection.execute(
                """DELETE FROM knowledge.personal_knowledge_documents
                   WHERE document_id = ? AND owner_user_id = ?""",
                (document_id, owner_user_id),
            )
            connection.execute(
                """UPDATE knowledge.personal_knowledge_jobs
                   SET status = 'deleted', stage = 'deleted', progress = 100, updated_at = ?
                   WHERE job_id = ? AND owner_user_id = ?""",
                (timestamp, job_id, owner_user_id),
            )
            connection.commit()
        emit("personal_knowledge_deleted", job_id, status="success", document_id=document_id)
    except Exception as exc:
        timestamp = _now()
        with db_session() as connection:
            connection.execute(
                """UPDATE knowledge.personal_knowledge_documents
                   SET status = 'failed', error_code = 'DELETE_FAILED',
                       error_message = '个人知识文档删除失败', updated_at = ?
                   WHERE document_id = ? AND owner_user_id = ?""",
                (timestamp, document_id, owner_user_id),
            )
            connection.execute(
                """UPDATE knowledge.personal_knowledge_jobs
                   SET status = 'failed', stage = 'failed', error_code = 'DELETE_FAILED',
                       error_message = '个人知识文档删除失败', retryable = 1, updated_at = ?
                   WHERE job_id = ? AND owner_user_id = ?""",
                (timestamp, job_id, owner_user_id),
            )
            connection.commit()
        emit("personal_knowledge_deleted", job_id, status="error", document_id=document_id, error_type=type(exc).__name__)


def personal_document_file(document_id: str, owner_user_id: str) -> tuple[Path, str, str]:
    with db_session() as connection:
        row = connection.execute(
            """SELECT filename, mime_type, storage_key, status FROM personal_knowledge_documents
               WHERE document_id = ? AND owner_user_id = ?""",
            (document_id, owner_user_id),
        ).fetchone()
    if row is None or row["status"] == "deleting":
        raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文档不存在")
    path = _storage_path(row["storage_key"])
    if not path.exists():
        raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识文件不存在")
    return path, row["filename"], row["mime_type"]


def _fts_query(text: str) -> str:
    terms = list(process_query_terms(text))
    for term in re.findall(r"[A-Za-z0-9_-]+|[\u4e00-\u9fff]{2,}", text):
        if re.fullmatch(r"[\u4e00-\u9fff]+", term) and len(term) > 4:
            terms.extend(term[index:index + 3] for index in range(len(term) - 2))
        else:
            terms.append(term)
    unique = list(dict.fromkeys(term for term in terms if len(term) >= 3))[:20]
    return " OR ".join(f'"{term}"' for term in unique)


def search_personal_knowledge(
    owner_user_id: str,
    query: str,
    limit: int = 10,
    document_ids: Iterable[str] = (),
) -> list[dict]:
    bounded = max(1, min(limit, 20))
    selected_ids = list(dict.fromkeys(document_ids))[:100]
    clauses = ["f.owner_user_id = ?", "d.status = 'ready'"]
    params: list[object] = [owner_user_id]
    if selected_ids:
        placeholders = ", ".join("?" for _ in selected_ids)
        with db_session() as connection:
            owned_count = connection.execute(
                f"""SELECT COUNT(*) FROM personal_knowledge_documents
                     WHERE owner_user_id = ? AND status = 'ready' AND document_id IN ({placeholders})""",
                [owner_user_id, *selected_ids],
            ).fetchone()[0]
        if owned_count != len(selected_ids):
            raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "指定文档不存在或尚未完成索引")
        clauses.append(f"f.document_id IN ({placeholders})")
        params.extend(selected_ids)
    fts_query = _fts_query(query)
    normalized_query = " ".join(query.casefold().split())
    tag_terms = list(dict.fromkeys([normalized_query, *(
        term.casefold() for term in re.findall(r"[A-Za-z0-9_-]+|[\u4e00-\u9fff]{2,}", query)
    )]))[:20]
    with db_session() as connection:
        # Exact element matching prevents a tag such as "sketch" from satisfying
        # "etch". Tag matches are intentionally placed before lexical FTS results.
        if tag_terms:
            tag_clauses = ["d.owner_user_id = ?", "d.status = 'ready'"]
            tag_params: list[object] = [owner_user_id]
            if selected_ids:
                placeholders = ", ".join("?" for _ in selected_ids)
                tag_clauses.append(f"d.document_id IN ({placeholders})")
                tag_params.extend(selected_ids)
            placeholders = ", ".join("?" for _ in tag_terms)
            tag_clauses.append(
                "EXISTS (SELECT 1 FROM json_each(d.tags_json) tag "
                f"WHERE LOWER(TRIM(tag.value)) IN ({placeholders}))"
            )
            tag_params.extend(tag_terms)
            tag_params.append(bounded)
            tag_rows = connection.execute(
                f"""SELECT c.chunk_id, c.document_id, d.title, c.section_path,
                           d.filename, d.process_steps_json, d.equipment_ids_json, d.tags_json, d.tags_source,
                           d.effective_date, c.page_number, c.content AS excerpt, -2.0 AS score
                    FROM personal_knowledge_documents d
                    JOIN personal_knowledge_chunks c ON c.chunk_id = (
                        SELECT c2.chunk_id FROM personal_knowledge_chunks c2
                        WHERE c2.document_id = d.document_id ORDER BY c2.chunk_index LIMIT 1
                    )
                    WHERE {' AND '.join(tag_clauses)}
                    ORDER BY d.updated_at DESC, d.document_id DESC LIMIT ?""",
                tag_params,
            ).fetchall()
        else:
            tag_rows = []
        if fts_query:
            rows = connection.execute(
                f"""SELECT f.chunk_id, f.document_id, f.title, f.section_path,
                           d.filename, d.process_steps_json, d.equipment_ids_json, d.tags_json, d.tags_source,
                           d.effective_date, c.page_number,
                           snippet(personal_knowledge_fts, 2, '[', ']', '…', 32) AS excerpt,
                           round(bm25(personal_knowledge_fts), 4) AS score
                    FROM knowledge.personal_knowledge_fts f
                    JOIN personal_knowledge_chunks c ON c.chunk_id = f.chunk_id
                    JOIN personal_knowledge_documents d ON d.document_id = f.document_id
                    WHERE personal_knowledge_fts MATCH ? AND {' AND '.join(clauses)}
                    ORDER BY bm25(personal_knowledge_fts) LIMIT ?""",
                [fts_query, *params, bounded * 2],
            ).fetchall()
        else:
            fallback_clauses = ["c.owner_user_id = ?", "d.status = 'ready'", "c.content LIKE ?"]
            fallback_params: list[object] = [owner_user_id, f"%{query.strip()}%"]
            if selected_ids:
                placeholders = ", ".join("?" for _ in selected_ids)
                fallback_clauses.append(f"c.document_id IN ({placeholders})")
                fallback_params.extend(selected_ids)
            fallback_params.append(bounded * 2)
            rows = connection.execute(
                f"""SELECT c.chunk_id, c.document_id, d.title, c.section_path,
                           d.filename, d.process_steps_json, d.equipment_ids_json, d.tags_json, d.tags_source,
                           d.effective_date, c.page_number, c.content AS excerpt, -1.0 AS score
                    FROM personal_knowledge_chunks c
                    JOIN personal_knowledge_documents d ON d.document_id = c.document_id
                    WHERE {' AND '.join(fallback_clauses)}
                    ORDER BY c.chunk_index LIMIT ?""",
                fallback_params,
            ).fetchall()
    seen_document_ids = {row["document_id"] for row in tag_rows}
    rows = list(tag_rows) + [row for row in rows if row["document_id"] not in seen_document_ids]
    rows = rows[:bounded]
    results = []
    for rank, row in enumerate(rows, start=1):
        item = dict(row)
        item["process_steps"] = _json_list(item.pop("process_steps_json", "[]"))
        item["equipment_ids"] = _json_list(item.pop("equipment_ids_json", "[]"))
        item["tags"] = _json_list(item.pop("tags_json", "[]"))
        item.update({
            "reference_id": f"PERSONAL-{item['chunk_id']}",
            "source_type": "personal",
            "source": "个人知识库",
            "section": item["section_path"],
            "retrieval": {
                "lexical_rank": rank,
                "final_score": item.pop("score"),
                "selection_reason": ["当前用户个人知识", f"lexical rank {rank}"],
            },
        })
        results.append(item)
    return results


def get_personal_chunk(chunk_id: str, owner_user_id: str) -> dict:
    with db_session() as connection:
        row = connection.execute(
            """SELECT c.chunk_id, c.document_id, d.title, d.filename, c.page_number,
                      c.section_path, c.chunk_index, c.content, c.token_count
               FROM personal_knowledge_chunks c
               JOIN personal_knowledge_documents d ON d.document_id = c.document_id
               WHERE c.chunk_id = ? AND c.owner_user_id = ? AND d.status = 'ready'""",
            (chunk_id, owner_user_id),
        ).fetchone()
    if row is None:
        raise PersonalKnowledgeError(404, "DOCUMENT_NOT_FOUND", "个人知识章节不存在")
    return dict(row)


def pending_personal_jobs() -> list[tuple[str, str, str, str]]:
    with db_session() as connection:
        rows = connection.execute(
            """SELECT j.document_id, j.owner_user_id, j.job_id, d.status
               FROM personal_knowledge_jobs j
               JOIN personal_knowledge_documents d ON d.document_id = j.document_id
               WHERE j.status IN ('uploaded', 'parsing', 'chunking', 'indexing', 'deleting')
                 AND j.created_at = (SELECT MAX(j2.created_at) FROM personal_knowledge_jobs j2
                                     WHERE j2.document_id = j.document_id)"""
        ).fetchall()
    return [(row["document_id"], row["owner_user_id"], row["job_id"], row["status"]) for row in rows]
