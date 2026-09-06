from __future__ import annotations

from pathlib import Path

from ..contracts import ExtractedSection


def extract_office(path: Path, suffix: str, title: str) -> list[ExtractedSection]:
    if suffix == ".docx":
        return _extract_docx(path, title)
    return _extract_xlsx(path)


def _extract_docx(path: Path, title: str) -> list[ExtractedSection]:
    from docx import Document

    document = Document(str(path))
    sections: list[ExtractedSection] = []
    heading = title
    body: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            if body:
                sections.append(ExtractedSection(None, heading, "\n".join(body)))
            heading = text
            body = []
        else:
            body.append(text)
    if body:
        sections.append(ExtractedSection(None, heading, "\n".join(body)))
    for index, table in enumerate(document.tables, start=1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        content = "\n".join(row for row in rows if row.strip()).strip()
        if content:
            sections.append(ExtractedSection(None, f"表格 {index}", content))
    return sections


def _extract_xlsx(path: Path) -> list[ExtractedSection]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[ExtractedSection] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    rows.append("\t".join(values).rstrip())
            if rows:
                sections.append(ExtractedSection(None, worksheet.title, "\n".join(rows)))
    finally:
        workbook.close()
    return sections
