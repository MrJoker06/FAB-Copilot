from __future__ import annotations

import atexit
import asyncio
import io
import json
import os
import shutil
import sqlite3
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx


TEST_DATABASE_DIR = Path(tempfile.mkdtemp(prefix="fab-insight-tests-"))
os.environ["FAB_KNOWLEDGE_DATABASE_PATH"] = str(TEST_DATABASE_DIR / "process_knowledge.db")
os.environ["FAB_AGENT_DATABASE_PATH"] = str(TEST_DATABASE_DIR / "ai_operations.db")
os.environ["FAB_PERSONAL_KNOWLEDGE_STORAGE_PATH"] = str(TEST_DATABASE_DIR / "personal_knowledge")
atexit.register(shutil.rmtree, TEST_DATABASE_DIR, True)

from fastapi.testclient import TestClient

from app.audit import read_events, sanitize
from app import agent as agent_module
from app.database import (
    db_session,
    department_paths,
    get_active_defect_aliases,
    initialize,
    replace_defect_type_aliases,
    sqlite_journal_mode,
)
from app.execution_store import (
    persist_execution_plan,
    recover_interrupted_executions,
    start_run_record,
    start_tool_event,
)
from app.main import app
from app.processes import MANUFACTURING_STEPS
from app.rag import _infer_filters, search_knowledge
from app.reports import process_capability, recent_excursions, report_bundle, safe_select, timestamp_bounds
from app.seed import seed_reference_data
from app.sessions import append_message, record_agent_result


class FabInsightTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        seed_reference_data()

    def test_dashboard_has_manufacturing_data(self):
        dashboard = report_bundle()
        self.assertGreater(dashboard["overview"]["lots"], 300)
        self.assertGreater(dashboard["overview"]["yield_pct"], 80)
        self.assertTrue(dashboard["equipment"])
        defect_types = [row["defect_type"] for row in dashboard["defects"]]
        self.assertEqual(len(defect_types), len(set(defect_types)))
        self.assertTrue(all("process_steps" in row for row in dashboard["defects"]))

    def test_health_exposes_runtime_instance_identity(self):
        with TestClient(app) as client:
            payload = client.get("/api/health").json()

        self.assertIn("instance", payload)
        self.assertTrue(payload["instance"]["id"])
        self.assertTrue(payload["instance"]["root"])
        self.assertTrue(payload["instance"]["source_root"])

    def test_frontend_exposes_personal_knowledge_upload_workflow(self):
        with TestClient(app) as client:
            response = client.get("/")
        self.assertEqual(response.status_code, 200)
        page = response.text
        for element_id in (
            "knowledgeSourceTabs",
            "openKnowledgeUpload",
            "knowledgeUploadDialog",
            "knowledgeUploadFile",
            "personalKnowledgeActions",
            "personalKnowledgeDeleteDialog",
        ):
            self.assertIn(f'id="{element_id}"', page)
        self.assertIn("/static/stream_buffer.js?v=20260719-1", page)
        self.assertIn("/static/source_presenter.js?v=20260719-1", page)
        self.assertIn("/static/app.js?v=20260719-1", page)
        self.assertNotIn('id="monitorView"', page)
        self.assertNotIn('id="apiKeyInput"', page)
        styles = (Path(__file__).parents[1] / "app" / "static" / "styles.css").read_text(encoding="utf-8")
        self.assertIn("scrollbar-gutter: stable", styles)
        script = (Path(__file__).parents[1] / "app" / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn("data-personal-chunk", script)
        self.assertIn("renderSourcePanel(messageSources)", script)
        self.assertIn("sourcePresenter.stripSourceAppendix(content)", script)
        self.assertIn("sourcePresenter.replaceReferenceIds(content, sources)", script)
        self.assertIn("event.key !== 'Enter' || event.shiftKey || event.isComposing || event.keyCode === 229", script)
        self.assertIn("$('#chatForm').requestSubmit()", script)
        self.assertIn("item.status === 'blocked'", script)
        self.assertIn("执行路径受限", script)
        self.assertIn("await finishLiveAnswer(live, result.answer)", script)
        self.assertNotIn("window.clearTimeout(live.renderTimer); live.renderTimer = null;", script)
        self.assertNotIn("item.arguments", script)
        self.assertNotIn("event.arguments", script)
        self.assertNotIn("argumentsText", script)
        self.assertIn(".trace-status-blocked", styles)
        self.assertIn(".source-list[open] summary::before", styles)

    def test_admin_frontend_is_separate_from_user_workspace(self):
        with TestClient(app) as client:
            response = client.get("/admin")
        self.assertEqual(response.status_code, 200)
        for element_id in (
            "adminAuthDialog",
            "runtimeView",
            "userStatusBody",
            "modelsView",
            "publishModelDialog",
        ):
            self.assertIn(f'id="{element_id}"', response.text)
        self.assertIn('<div id="adminWorkspace" class="admin-workspace" hidden>', response.text)
        self.assertIn("/static/admin.js?v=20260718-3", response.text)
        script = (Path(__file__).parents[1] / "app" / "static" / "admin.js").read_text(encoding="utf-8")
        self.assertIn("/api/admin/overview", script)
        self.assertIn("/api/admin/models", script)
        self.assertIn("$('#adminWorkspace').hidden = true", script)
        self.assertIn("$('#adminWorkspace').hidden = false", script)
        self.assertIn("run.question_preview", script)
        self.assertIn("run.session_id", script)
        self.assertIn("问题 / Session", response.text)

    def test_seeded_cvd_drift_is_detectable(self):
        capability = process_capability(process_step="CVD")
        cvd_02 = next(row for row in capability if row["equipment_id"] == "CVD-02")
        peers = [row for row in capability if row["equipment_id"] != "CVD-02"]
        self.assertLess(cvd_02["cpk"], min(row["cpk"] for row in peers))

    def test_capability_range_is_sargable_and_keeps_end_date_inclusive(self):
        self.assertEqual(timestamp_bounds("2026-07-01", "2026-07-13"), ("2026-07-01", "2026-07-14"))
        with db_session() as connection:
            plan = connection.execute(
                "EXPLAIN QUERY PLAN SELECT equipment_id FROM process_runs "
                "WHERE process_step = ? AND run_time >= ? AND run_time < ?",
                ("CMP", "2026-07-01", "2026-07-14"),
            ).fetchall()
        self.assertTrue(any("idx_runs_step_time" in row[3] for row in plan), plan)

    def test_all_manufacturing_steps_have_equipment_and_capability(self):
        with db_session() as connection:
            rows = connection.execute(
                "SELECT process_step, COUNT(DISTINCT equipment_id) AS equipment_count "
                "FROM process_runs GROUP BY process_step"
            ).fetchall()
        equipment_counts = {row["process_step"]: row["equipment_count"] for row in rows}
        self.assertEqual(set(equipment_counts), set(MANUFACTURING_STEPS))
        self.assertTrue(all(count >= 3 for count in equipment_counts.values()))
        for step in MANUFACTURING_STEPS:
            with self.subTest(step=step):
                capability = process_capability(process_step=step)
                self.assertEqual(len(capability), equipment_counts[step])
                self.assertTrue(all(row["metric_name"] and row["metric_unit"] for row in capability))

    def test_expanded_fab_history_has_route_sensor_and_hidden_ground_truth(self):
        with db_session() as connection:
            counts = {
                table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                for table in ("lots", "process_runs", "sensor_readings", "equipment_state_intervals", "wafer_maps")
            }
            scenarios = connection.execute(
                "SELECT COUNT(*), SUM(visible_to_agent) FROM simulation_scenarios"
            ).fetchone()
            max_operation = connection.execute("SELECT MAX(operation_seq) FROM process_runs").fetchone()[0]
            signal_coverage = connection.execute(
                "SELECT process_step, COUNT(DISTINCT signal_name) FROM sensor_readings GROUP BY process_step"
            ).fetchall()
        self.assertGreaterEqual(counts["lots"], 9000)
        self.assertGreaterEqual(counts["process_runs"], 150000)
        self.assertGreaterEqual(counts["sensor_readings"], 1000000)
        self.assertGreater(counts["equipment_state_intervals"], 25000)
        self.assertEqual(counts["wafer_maps"], counts["lots"])
        self.assertEqual(tuple(scenarios), (9, 0))
        self.assertEqual(max_operation, 18)
        self.assertEqual({row[0] for row in signal_coverage}, set(MANUFACTURING_STEPS))
        self.assertTrue(all(row[1] >= 6 for row in signal_coverage))

    def test_history_contains_no_future_process_records(self):
        with db_session() as connection:
            max_run = connection.execute("SELECT MAX(date(run_time)) FROM process_runs").fetchone()[0]
            max_defect = connection.execute("SELECT MAX(date(detected_time)) FROM defect_records").fetchone()[0]
            max_wafer_map = connection.execute("SELECT MAX(date(detected_time)) FROM wafer_maps").fetchone()[0]
        self.assertLessEqual(max_run, date.today().isoformat())
        self.assertLessEqual(max_defect, date.today().isoformat())
        self.assertLessEqual(max_wafer_map, date.today().isoformat())

    def test_department_databases_are_physically_separated(self):
        paths = department_paths()
        self.assertEqual(set(paths), {"main", "equipment", "quality", "knowledge", "agent"})
        self.assertEqual(len(set(paths.values())), 5)
        self.assertTrue(all(path.exists() for path in paths.values()))
        with db_session() as connection:
            attached = {row[1] for row in connection.execute("PRAGMA database_list")}
            plan_columns = {row[1] for row in connection.execute("PRAGMA agent.table_info(agent_execution_plans)")}
            step_columns = {row[1] for row in connection.execute("PRAGMA agent.table_info(agent_plan_steps)")}
            tool_columns = {row[1] for row in connection.execute("PRAGMA agent.table_info(agent_tool_events)")}
        self.assertTrue({"main", "equipment", "quality", "knowledge", "agent"}.issubset(attached))
        self.assertTrue({"plan_id", "definition_status", "execution_status"}.issubset(plan_columns))
        self.assertTrue({"step_id", "status", "attempt_count"}.issubset(step_columns))
        self.assertTrue({"plan_step_id", "started_at", "completed_at"}.issubset(tool_columns))

    def test_sqlite_journal_mode_is_safe_for_wsl_mounted_databases(self):
        with patch.dict(os.environ, {"FAB_SQLITE_JOURNAL_MODE": "auto"}):
            self.assertEqual(sqlite_journal_mode(Path("/mnt/d/fab.db")), "DELETE")
            self.assertEqual(sqlite_journal_mode(Path("/tmp/fab.db")), "WAL")
        with patch.dict(os.environ, {"FAB_SQLITE_JOURNAL_MODE": "TRUNCATE"}):
            self.assertEqual(sqlite_journal_mode(Path("/mnt/d/fab.db")), "TRUNCATE")

    def test_agent_run_migration_backfills_question_and_user_message(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "legacy-agent.db"
            connection = sqlite3.connect(database_path)
            connection.executescript(
                """CREATE TABLE chat_sessions (
                       session_id TEXT PRIMARY KEY, tenant_id TEXT NOT NULL DEFAULT 'demo-fab',
                       owner_user_id TEXT NOT NULL DEFAULT 'local-user', title TEXT NOT NULL,
                       created_at TEXT NOT NULL, updated_at TEXT NOT NULL, model TEXT NOT NULL DEFAULT '',
                       archived_at TEXT, last_message_preview TEXT NOT NULL DEFAULT ''
                   );
                   CREATE TABLE chat_messages (
                       message_id INTEGER PRIMARY KEY AUTOINCREMENT,
                       session_id TEXT NOT NULL REFERENCES chat_sessions(session_id) ON DELETE CASCADE,
                       role TEXT NOT NULL, content TEXT NOT NULL, metadata_json TEXT NOT NULL DEFAULT '{}',
                       client_message_id TEXT, created_at TEXT NOT NULL
                   );
                   CREATE TABLE agent_runs (
                       run_id TEXT PRIMARY KEY,
                       session_id TEXT NOT NULL REFERENCES chat_sessions(session_id) ON DELETE CASCADE,
                       assistant_message_id INTEGER REFERENCES chat_messages(message_id) ON DELETE SET NULL,
                       model TEXT NOT NULL, mode TEXT NOT NULL, status TEXT NOT NULL,
                       started_at TEXT NOT NULL, completed_at TEXT NOT NULL,
                       first_token_ms REAL, duration_ms REAL,
                       input_tokens INTEGER NOT NULL DEFAULT 0, output_tokens INTEGER NOT NULL DEFAULT 0,
                       error_message TEXT
                   );
                   INSERT INTO chat_sessions(session_id, title, created_at, updated_at)
                   VALUES ('legacy-session', '用户修改后的名称', '2026-07-17', '2026-07-17');
                   INSERT INTO chat_messages(session_id, role, content, created_at)
                   VALUES ('legacy-session', 'user', '分析 ETCH-04 最近三天故障', '2026-07-17T01:00:00Z');
                   INSERT INTO chat_messages(session_id, role, content, created_at)
                   VALUES ('legacy-session', 'assistant', '分析结果', '2026-07-17T01:01:00Z');
                   INSERT INTO agent_runs(
                       run_id, session_id, assistant_message_id, model, mode, status, started_at, completed_at
                   ) VALUES ('legacy-run', 'legacy-session', 2, 'test-model', 'test', 'success',
                             '2026-07-17T01:00:00Z', '2026-07-17T01:01:00Z');"""
            )
            connection.commit()
            connection.close()

            initialize(database_path)
            connection = sqlite3.connect(database_path)
            row = connection.execute(
                "SELECT user_message_id, question_preview FROM agent_runs WHERE run_id = 'legacy-run'"
            ).fetchone()
            columns = {item[1] for item in connection.execute("PRAGMA table_info(agent_runs)")}
            connection.close()

        self.assertEqual(row, (1, "分析 ETCH-04 最近三天故障"))
        self.assertTrue({"user_message_id", "question_preview"}.issubset(columns))

    def test_recent_excursions_supports_process_filter(self):
        rows = recent_excursions(process_step="CMP", limit=20)
        self.assertTrue(rows)
        self.assertTrue(all(row["process_step"] == "CMP" for row in rows))

    def test_rag_returns_cited_source(self):
        results = search_knowledge("CMP 划伤 Pad Conditioner", 10)
        self.assertTrue(results)
        self.assertEqual({item["source_type"] for item in results}, {"knowledge", "case"})
        self.assertTrue(all(item["reference_id"] for item in results))
        self.assertTrue(all(item["retrieval"]["selection_reason"] for item in results))

    def test_rag_infers_chinese_process_and_defect_aliases(self):
        results = search_knowledge("刻蚀设备 CD 偏移的潜在原因", 6)
        self.assertTrue(results)
        self.assertEqual({item["source_type"] for item in results}, {"knowledge", "case"})
        self.assertTrue(all(item["retrieval"]["filters"]["process_step"] == "ETCH" for item in results))
        self.assertTrue(all(item["retrieval"]["filters"]["defect_type"] == "CD Shift" for item in results))

    def test_rag_uses_database_managed_defect_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "aliases.db"
            initialize(database_path)
            replace_defect_type_aliases("Pattern Collapse", ("pattern collapse", "图形倒塌"), database_path)
            aliases = get_active_defect_aliases(database_path)
        with patch("app.rag.get_active_defect_aliases", return_value=aliases):
            filters = _infer_filters("光刻后出现图形倒塌", None, None, None)
        self.assertEqual(filters["defect_type"], "Pattern Collapse")

    def test_rag_gold_dataset(self):
        dataset_path = Path(__file__).resolve().parent.parent / "evaluations" / "rag_gold.json"
        dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
        for item in dataset:
            with self.subTest(item=item["id"]):
                results = search_knowledge(item["query"], limit=6, **item.get("filters", {}))
                references = {result["reference_id"] for result in results}
                source_types = {result["source_type"] for result in results}
                self.assertTrue(references.intersection(item["expected_any"]))
                self.assertTrue(set(item.get("required_source_types", [])).issubset(source_types))
                self.assertFalse(references.intersection(item.get("forbidden", [])))

    def test_knowledge_documents_support_chunk_navigation(self):
        with TestClient(app) as client:
            documents = client.get("/api/knowledge/documents")
            self.assertEqual(documents.status_code, 200)
            self.assertGreaterEqual(len(documents.json()), 12)
            first_chunk_id = documents.json()[0]["first_chunk_id"]
            document = client.get(f"/api/knowledge/chunks/{first_chunk_id}")
            self.assertEqual(document.status_code, 200)
            self.assertEqual(document.json()["selected_chunk_id"], first_chunk_id)
            self.assertTrue(document.json()["chunks"])
            self.assertTrue(all("content" in chunk for chunk in document.json()["chunks"]))
            self.assertEqual(document.json()["document"]["status"], "Active")
            self.assertTrue(document.json()["document"]["version"])
            self.assertGreaterEqual(document.json()["document"]["authority"], 1)

    def test_personal_knowledge_upload_search_privacy_and_delete(self):
        content = (
            "# CMP-03 Pad 更换记录\n\n"
            "2026-07-01 完成 Pad 更换。更换后确认主轴电流稳定，抛光压力恢复正常。\n\n"
            "## 验证结果\n\n"
            "连续三个批次膜厚均匀性满足要求，Scratch 缺陷未再次出现。"
        ).encode("utf-8")
        headers = {"Idempotency-Key": "personal-upload-test-001"}
        metadata = json.dumps({
            "title": "CMP-03 Pad 更换记录",
            "process_steps": ["CMP"],
            "equipment_ids": ["CMP-03"],
            "effective_date": "2026-07-01",
            "tags": ["pad", "maintenance"],
        }, ensure_ascii=False)

        with TestClient(app) as client:
            uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers=headers,
                files={"file": ("cmp-pad.md", content, "text/markdown")},
                data={"metadata": metadata},
            )
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            document_id = uploaded.json()["document_id"]
            job_id = uploaded.json()["job_id"]

            job = client.get(f"/api/knowledge/jobs/{job_id}")
            self.assertEqual(job.status_code, 200)
            self.assertEqual(job.json()["status"], "ready")
            self.assertEqual(job.json()["progress"], 100)

            listed = client.get("/api/knowledge/personal/documents?status=ready&q=CMP")
            self.assertEqual(listed.status_code, 200)
            self.assertIn(document_id, {item["document_id"] for item in listed.json()["items"]})

            detail = client.get(f"/api/knowledge/personal/documents/{document_id}")
            self.assertEqual(detail.status_code, 200)
            self.assertEqual(detail.json()["process_steps"], ["CMP"])
            self.assertGreater(detail.json()["chunk_count"], 0)
            self.assertNotIn("owner_user_id", detail.json())
            self.assertNotIn("storage_key", detail.json())

            document_content = client.get(f"/api/knowledge/personal/documents/{document_id}/content")
            self.assertEqual(document_content.status_code, 200)
            self.assertTrue(document_content.json()["chunks"])
            chunk_id = document_content.json()["chunks"][0]["chunk_id"]
            with db_session() as connection:
                indexed_chunks = connection.execute(
                    "SELECT COUNT(*) FROM personal_knowledge_fts WHERE document_id = ?", (document_id,)
                ).fetchone()[0]
            self.assertEqual(indexed_chunks, detail.json()["chunk_count"])

            search = client.post(
                "/api/knowledge/personal/search-preview",
                json={"query": "CMP-03 Pad 更换后 Scratch 验证", "document_ids": [document_id]},
            )
            self.assertEqual(search.status_code, 200, search.text)
            self.assertGreater(search.json()["count"], 0)
            self.assertTrue(all(item["source_type"] == "personal" for item in search.json()["results"]))

            chunk = client.get(f"/api/knowledge/personal/chunks/{chunk_id}")
            self.assertEqual(chunk.status_code, 200)
            self.assertEqual(chunk.json()["document_id"], document_id)

            updated = client.patch(
                f"/api/knowledge/personal/documents/{document_id}",
                json={"title": "CMP-03 Pad 更换与验证记录", "tags": ["pad", "verified"]},
            )
            self.assertEqual(updated.status_code, 200, updated.text)
            self.assertEqual(updated.json()["title"], "CMP-03 Pad 更换与验证记录")

            downloaded = client.get(f"/api/knowledge/personal/documents/{document_id}/file")
            self.assertEqual(downloaded.status_code, 200)
            self.assertEqual(downloaded.content, content)

            replay = client.post(
                "/api/knowledge/personal/documents",
                headers=headers,
                files={"file": ("cmp-pad.md", content, "text/markdown")},
                data={"metadata": metadata},
            )
            self.assertEqual(replay.status_code, 202)
            self.assertEqual(replay.json()["document_id"], document_id)
            self.assertEqual(replay.json()["job_id"], job_id)

            duplicate = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "personal-upload-test-002"},
                files={"file": ("duplicate.md", content, "text/markdown")},
            )
            self.assertEqual(duplicate.status_code, 409)
            self.assertEqual(duplicate.json()["error_code"], "DUPLICATE_DOCUMENT")

            with patch.dict(os.environ, {"FAB_PERSONAL_KNOWLEDGE_USER_ID": "different-user"}):
                self.assertEqual(client.get(f"/api/knowledge/personal/documents/{document_id}").status_code, 404)
                self.assertEqual(client.get(f"/api/knowledge/personal/chunks/{chunk_id}").status_code, 404)
                self.assertEqual(client.get(f"/api/knowledge/jobs/{job_id}").status_code, 404)
                private_search = client.post(
                    "/api/knowledge/personal/search-preview",
                    json={"query": "CMP Pad", "document_ids": [document_id]},
                )
                self.assertEqual(private_search.status_code, 404)

            deleting = client.delete(f"/api/knowledge/personal/documents/{document_id}")
            self.assertEqual(deleting.status_code, 202, deleting.text)
            delete_job_id = deleting.json()["job_id"]
            self.assertEqual(client.get(f"/api/knowledge/jobs/{delete_job_id}").json()["status"], "deleted")
            self.assertEqual(client.get(f"/api/knowledge/personal/documents/{document_id}").status_code, 404)
            self.assertEqual(client.get(f"/api/knowledge/personal/chunks/{chunk_id}").status_code, 404)
            with db_session() as connection:
                indexed_chunks = connection.execute(
                    "SELECT COUNT(*) FROM personal_knowledge_fts WHERE document_id = ?", (document_id,)
                ).fetchone()[0]
            self.assertEqual(indexed_chunks, 0)

    def test_agent_search_combines_private_knowledge_with_cross_language_recall(self):
        content = (
            "# Introduction to Photolithography\n\n"
            "Photolithography transfers a reticle pattern into photoresist by optical exposure.\n\n"
            "## Process basics\n\n"
            "The basic sequence covers coat, soft bake, align, expose, develop, inspect, and pattern transfer."
        ).encode("utf-8")
        metadata = json.dumps({
            "title": "Introduction to Photolithography",
            "process_steps": ["LITHO"],
            "tags": ["training", "fundamentals"],
        })
        with TestClient(app) as client:
            uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "agent-personal-rag-test-001"},
                files={"file": ("intro-photolithography.md", content, "text/markdown")},
                data={"metadata": metadata},
            )
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            document_id = uploaded.json()["document_id"]
            try:
                results = agent_module.search_agent_knowledge(
                    "请介绍光刻技术原理",
                    limit=6,
                    source_types=["personal"],
                )
                self.assertTrue(results)
                self.assertTrue(all(item["source_type"] == "personal" for item in results))
                self.assertEqual(results[0]["document_id"], document_id)
                self.assertEqual(results[0]["source"], "个人知识库")
                self.assertEqual(results[0]["process_steps"], ["LITHO"])
                self.assertTrue(results[0]["reference_id"].startswith("PERSONAL-"))
                self.assertNotIn(
                    "personal",
                    {item["source_type"] for item in search_knowledge("请介绍光刻技术原理", limit=6)},
                )
            finally:
                deleted = client.delete(f"/api/knowledge/personal/documents/{document_id}")
                self.assertEqual(deleted.status_code, 202, deleted.text)

    def test_personal_knowledge_upload_validation_and_token(self):
        with TestClient(app) as client:
            missing_key = client.post(
                "/api/knowledge/personal/documents",
                files={"file": ("note.txt", b"valid text content", "text/plain")},
            )
            self.assertEqual(missing_key.status_code, 422)

            unsupported = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "unsupported-file-test"},
                files={"file": ("payload.exe", b"not executable", "application/octet-stream")},
            )
            self.assertEqual(unsupported.status_code, 415)
            self.assertEqual(unsupported.json()["error_code"], "UNSUPPORTED_FILE_TYPE")

            with patch.dict(os.environ, {"FAB_PERSONAL_KNOWLEDGE_API_TOKEN": "private-test-token"}):
                unauthenticated = client.get("/api/knowledge/personal/documents")
                self.assertEqual(unauthenticated.status_code, 401)
                authenticated = client.get(
                    "/api/knowledge/personal/documents",
                    headers={"Authorization": "Bearer private-test-token"},
                )
                self.assertEqual(authenticated.status_code, 200)

    def test_personal_knowledge_scanned_pdf_requires_ocr(self):
        from pypdf import PdfWriter

        output = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(output)
        with TestClient(app) as client:
            uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "ocr-required-test"},
                files={"file": ("scanned.pdf", output.getvalue(), "application/pdf")},
            )
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            job = client.get(f"/api/knowledge/jobs/{uploaded.json()['job_id']}")
            self.assertEqual(job.status_code, 200)
            self.assertEqual(job.json()["status"], "failed")
            self.assertEqual(job.json()["error"]["error_code"], "OCR_REQUIRED")
            deleted = client.delete(f"/api/knowledge/personal/documents/{uploaded.json()['document_id']}")
            self.assertEqual(deleted.status_code, 202)

    def test_personal_knowledge_office_and_csv_parsers(self):
        from docx import Document
        from openpyxl import Workbook

        docx_buffer = io.BytesIO()
        docx_document = Document()
        docx_document.add_heading("ETCH-01 维护记录", level=1)
        docx_document.add_paragraph("腔体清洗完成，Particle 复测符合要求。")
        docx_document.save(docx_buffer)

        xlsx_buffer = io.BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CVD-02 检查"
        sheet.append(["项目", "结果"])
        sheet.append(["膜厚均匀性", "通过"])
        workbook.save(xlsx_buffer)
        workbook.close()

        samples = (
            ("etch-maintenance.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "腔体清洗"),
            ("cvd-check.xlsx", xlsx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "膜厚均匀性"),
            ("cmp-check.csv", "项目,结果\nPad检查,通过\n".encode("utf-8"), "text/csv", "Pad检查"),
        )
        with TestClient(app) as client:
            for index, (filename, content, mime_type, expected) in enumerate(samples):
                with self.subTest(filename=filename):
                    uploaded = client.post(
                        "/api/knowledge/personal/documents",
                        headers={"Idempotency-Key": f"parser-format-test-{index}"},
                        files={"file": (filename, content, mime_type)},
                    )
                    self.assertEqual(uploaded.status_code, 202, uploaded.text)
                    document_id = uploaded.json()["document_id"]
                    job = client.get(f"/api/knowledge/jobs/{uploaded.json()['job_id']}")
                    self.assertEqual(job.json()["status"], "ready", job.text)
                    parsed = client.get(f"/api/knowledge/personal/documents/{document_id}/content")
                    self.assertIn(expected, " ".join(chunk["content"] for chunk in parsed.json()["chunks"]))
                    self.assertEqual(client.delete(f"/api/knowledge/personal/documents/{document_id}").status_code, 202)

    def test_historical_case_has_fixed_format(self):
        with TestClient(app) as client:
            cases = client.get("/api/cases")
            self.assertEqual(cases.status_code, 200)
            self.assertGreaterEqual(len(cases.json()), 12)
            case_id = cases.json()[0]["case_id"]
            detail = client.get(f"/api/cases/{case_id}")
            self.assertEqual(detail.status_code, 200)
            required = {
                "problem", "solution", "result", "summary", "root_cause",
                "root_cause_confidence", "validation_method", "recurrence_status",
                "approved_by", "case_status",
            }
            self.assertTrue(required.issubset(detail.json()))
            self.assertTrue(all(detail.json()[field] for field in required))

    def test_sql_guard_rejects_writes(self):
        with self.assertRaises(ValueError):
            safe_select("DELETE FROM lots")
        with self.assertRaises(ValueError):
            safe_select("SELECT * FROM sqlite_master")
        with self.assertRaises(ValueError):
            safe_select("SELECT scenario_id FROM process_runs")

    def test_health_endpoint(self):
        with TestClient(app) as client:
            response = client.get("/api/health")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["status"], "ok")
            self.assertEqual(response.json()["reference_data"]["knowledge_chunks"], 68)
            self.assertEqual(response.json()["reference_data"]["historical_cases"], 17)
            self.assertTrue(response.json()["build"]["version"])
            self.assertTrue(response.json()["enterprise"]["simulated"])
            self.assertEqual(response.json()["enterprise"]["virtual_scale"]["lots"], 25_000_000)

    def test_enterprise_context_resolves_request_identity(self):
        headers = {"X-Tenant-ID": "fab-east", "X-User-ID": "engineer-07", "X-Roles": "analyst,process-engineer"}
        with TestClient(app) as client:
            response = client.get("/api/enterprise/context", headers=headers)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["identity"]["tenant_id"], "fab-east")
            self.assertEqual(response.json()["identity"]["user_id"], "engineer-07")
            self.assertEqual(response.json()["knowledge_owner_scope"], "fab-east:engineer-07")
            self.assertEqual(response.json()["identity"]["roles"], ["analyst", "process-engineer"])

            invalid = client.get(
                "/api/enterprise/context",
                headers={"X-Tenant-ID": "invalid tenant", "X-User-ID": "engineer-07"},
            )
            self.assertEqual(invalid.status_code, 400)
            self.assertEqual(invalid.json()["error_code"], "INVALID_IDENTITY")

    def test_sessions_are_isolated_by_tenant_and_user(self):
        owner_headers = {"X-Tenant-ID": "fab-a", "X-User-ID": "owner"}
        other_tenant_headers = {"X-Tenant-ID": "fab-b", "X-User-ID": "owner"}
        other_user_headers = {"X-Tenant-ID": "fab-a", "X-User-ID": "other"}
        with TestClient(app) as client:
            created = client.post("/api/sessions", headers=owner_headers)
            self.assertEqual(created.status_code, 200)
            session_id = created.json()["session_id"]

            self.assertEqual(client.get(f"/api/sessions/{session_id}", headers=owner_headers).status_code, 200)
            self.assertEqual(client.get(f"/api/sessions/{session_id}", headers=other_tenant_headers).status_code, 404)
            self.assertEqual(client.get(f"/api/sessions/{session_id}", headers=other_user_headers).status_code, 404)
            self.assertNotIn(session_id, {item["session_id"] for item in client.get("/api/sessions", headers=other_tenant_headers).json()})
            self.assertEqual(client.delete(f"/api/sessions/{session_id}/permanent", headers=owner_headers).status_code, 200)

    def test_enterprise_mock_data_is_deterministic_and_cursor_scoped(self):
        headers = {"X-Tenant-ID": "fab-a", "X-User-ID": "analyst"}
        path = "/api/enterprise/data/process-runs?process_step=ETCH&limit=3"
        with TestClient(app) as client:
            first = client.get(path, headers=headers)
            repeated = client.get(path, headers=headers)
            self.assertEqual(first.status_code, 200, first.text)
            self.assertEqual(first.json()["records"], repeated.json()["records"])
            self.assertTrue(all(item["process_step"] == "ETCH" and item["synthetic"] for item in first.json()["records"]))
            cursor = first.json()["page"]["next_cursor"]
            self.assertTrue(cursor)

            second = client.get(
                "/api/enterprise/data/process-runs",
                headers=headers,
                params={"process_step": "ETCH", "limit": 3, "cursor": cursor},
            )
            self.assertEqual(second.status_code, 200, second.text)
            self.assertFalse(
                {item["run_id"] for item in first.json()["records"]}
                & {item["run_id"] for item in second.json()["records"]}
            )

            cross_tenant = client.get(
                "/api/enterprise/data/process-runs",
                headers={"X-Tenant-ID": "fab-b", "X-User-ID": "analyst"},
                params={"process_step": "ETCH", "limit": 3, "cursor": cursor},
            )
            self.assertEqual(cross_tenant.status_code, 400)
            self.assertEqual(cross_tenant.json()["error_code"], "CURSOR_SCOPE_MISMATCH")

            midpoint = len(cursor) // 2
            replacement = "A" if cursor[midpoint] != "A" else "B"
            tampered = f"{cursor[:midpoint]}{replacement}{cursor[midpoint + 1:]}"
            invalid = client.get(
                "/api/enterprise/data/process-runs",
                headers=headers,
                params={"process_step": "ETCH", "limit": 3, "cursor": tampered},
            )
            self.assertEqual(invalid.status_code, 400)
            self.assertEqual(invalid.json()["error_code"], "INVALID_CURSOR")

    def test_enterprise_summary_is_projected_deterministic_and_tenant_scoped(self):
        end = date.today()
        start = end - timedelta(days=6)
        params = {"start": start.isoformat(), "end": end.isoformat(), "process_step": "CMP"}
        tenant_a = {"X-Tenant-ID": "fab-summary-a", "X-User-ID": "analyst"}
        tenant_b = {"X-Tenant-ID": "fab-summary-b", "X-User-ID": "analyst"}
        with TestClient(app) as client:
            first = client.get("/api/enterprise/data/summary", headers=tenant_a, params=params)
            repeated = client.get("/api/enterprise/data/summary", headers=tenant_a, params=params)
            self.assertEqual(first.status_code, 200, first.text)
            self.assertEqual(first.json(), repeated.json())
            payload = first.json()
            self.assertTrue(payload["simulated"])
            self.assertFalse(payload["aggregation"]["exact"])
            self.assertEqual(payload["aggregation"]["semantics"], "deterministic_projection")
            self.assertEqual(payload["query"]["filters"]["process_step"], "CMP")
            self.assertGreater(payload["totals"]["lots"], 0)
            self.assertGreater(payload["totals"]["process_runs"], payload["totals"]["lots"])
            self.assertGreater(payload["totals"]["sensor_events"], payload["totals"]["process_runs"])
            self.assertEqual(len(payload["equipment"]), 5)
            self.assertEqual(
                [item["health_score"] for item in payload["equipment"]],
                sorted(item["health_score"] for item in payload["equipment"]),
            )
            self.assertEqual(payload["lineage"]["classification"], "synthetic_non_production")

            other_tenant = client.get("/api/enterprise/data/summary", headers=tenant_b, params=params)
            self.assertEqual(other_tenant.status_code, 200)
            for field in ("lots", "wafers", "process_runs", "sensor_events"):
                self.assertEqual(other_tenant.json()["totals"][field], payload["totals"][field])
            self.assertNotEqual(other_tenant.json()["query"]["query_id"], payload["query"]["query_id"])
            self.assertNotEqual(other_tenant.json()["kpis"], payload["kpis"])

            equipment = client.get(
                "/api/enterprise/data/summary",
                headers=tenant_a,
                params={**params, "equipment_id": "CMP-03"},
            )
            self.assertEqual(equipment.status_code, 200)
            self.assertEqual([item["equipment_id"] for item in equipment.json()["equipment"]], ["CMP-03"])
            self.assertLess(equipment.json()["totals"]["lots"], payload["totals"]["lots"])

            invalid = client.get(
                "/api/enterprise/data/summary",
                headers=tenant_a,
                params={"start": end.isoformat(), "end": start.isoformat()},
            )
            self.assertEqual(invalid.status_code, 400)
            self.assertEqual(invalid.json()["error_code"], "INVALID_DATE_RANGE")

    def test_api_errors_include_correlation_id_and_retry_metadata(self):
        request_id = "req-test-correlation-001"
        with TestClient(app) as client:
            response = client.get("/api/sessions/missing-session", headers={"X-Request-ID": request_id})
            self.assertEqual(response.status_code, 404)
            self.assertEqual(response.headers["x-request-id"], request_id)
            self.assertEqual(response.json()["request_id"], request_id)
            self.assertEqual(response.json()["error_code"], "NOT_FOUND")
            self.assertFalse(response.json()["retryable"])

            invalid = client.post("/api/chat/stream", json={"message": "x"})
            self.assertEqual(invalid.status_code, 422)
            self.assertEqual(invalid.json()["error_code"], "VALIDATION_ERROR")
            self.assertTrue(invalid.json()["request_id"])

    def test_unexpected_500_is_structured_without_exposing_traceback(self):
        request_id = "req-internal-error-001"
        with patch("app.main.report_bundle", side_effect=RuntimeError("database secret detail")), TestClient(app, raise_server_exceptions=False) as client:
            response = client.get("/api/reports/dashboard", headers={"X-Request-ID": request_id})
            self.assertEqual(response.status_code, 500)
            self.assertEqual(response.headers["x-request-id"], request_id)
            self.assertEqual(response.json()["error_code"], "INTERNAL_ERROR")
            self.assertEqual(response.json()["request_id"], request_id)
            self.assertTrue(response.json()["retryable"])
            self.assertNotIn("database secret detail", response.text)

    def test_monitor_endpoint_is_sanitized(self):
        with patch.dict(os.environ, {"FAB_ADMIN_API_TOKEN": "monitor-admin-token"}), TestClient(app) as client:
            self.assertEqual(client.get("/api/monitor/summary?hours=24").status_code, 401)
            response = client.get(
                "/api/monitor/summary?hours=24",
                headers={"Authorization": "Bearer monitor-admin-token"},
            )
            self.assertEqual(response.status_code, 200)
            payload = response.json()
            self.assertIn("runs", payload)
            self.assertNotIn("api_key", str(payload).lower())

    def test_model_can_only_be_selected_from_published_catalog(self):
        with TestClient(app) as client:
            initial = client.get("/api/settings/model")
            self.assertEqual(initial.status_code, 200)
            self.assertEqual(initial.json()["allowed_models"], ["deepseekv4-flash", "deepseekv4-pro"])
            self.assertEqual(len(initial.json()["model_options"]), 2)
            self.assertEqual(initial.json()["credential_scope"], "server_managed")
            self.assertNotIn("api_key_masked", initial.json())
            self.assertNotIn("api_key_source", initial.json())
            response = client.post("/api/settings/model", json={"model": "deepseekv4-pro"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["model"], "deepseekv4-pro")
            self.assertEqual(response.json()["provider_model"], "deepseek-v4-pro")
            rejected = client.post("/api/settings/model", json={
                "model": "deepseekv4-pro",
                "api_key": "user-supplied-key-must-be-rejected",
            })
            self.assertEqual(rejected.status_code, 422, rejected.text)
            self.assertNotIn("user-supplied-key-must-be-rejected", rejected.text)
            client.post("/api/settings/model", json={"model": "deepseekv4-flash"})

    def test_admin_can_publish_model_and_user_selection_is_scoped(self):
        model_id = "enterprise-test-model"
        admin_token = "enterprise-admin-token"
        with db_session() as connection:
            connection.execute("DELETE FROM agent.user_model_preferences WHERE model_id = ?", (model_id,))
            connection.execute("DELETE FROM agent.model_catalog WHERE model_id = ?", (model_id,))
            connection.commit()
        try:
            with patch.dict(os.environ, {
                "FAB_ADMIN_API_TOKEN": admin_token,
                "DEEPSEEK_API_KEY": "server-managed-test-key",
            }), TestClient(app) as client:
                self.assertEqual(client.get("/api/admin/models").status_code, 401)
                admin_headers = {"Authorization": f"Bearer {admin_token}"}
                published = client.post(
                    "/api/admin/models",
                    headers=admin_headers,
                    json={
                        "model_id": model_id,
                        "display_name": "Enterprise Test Model",
                        "provider": "openai-compatible",
                        "provider_model": "enterprise-test-v1",
                        "base_url": "https://models.example.test/v1",
                        "capabilities": {"streaming": True, "tool_calls": True},
                    },
                )
                self.assertEqual(published.status_code, 201, published.text)
                self.assertEqual(published.json()["status"], "published")
                invalid = client.post(
                    "/api/admin/models",
                    headers=admin_headers,
                    json={
                        "model_id": "invalid-blank-name",
                        "display_name": "  ",
                        "provider": "openai-compatible",
                        "provider_model": "enterprise-test-v1",
                        "base_url": "https://models.example.test/v1",
                    },
                )
                self.assertEqual(invalid.status_code, 400, invalid.text)

                user_a = {"X-Tenant-ID": "model-fab", "X-User-ID": "user-a"}
                user_b = {"X-Tenant-ID": "model-fab", "X-User-ID": "user-b"}
                catalog = client.get("/api/settings/model", headers=user_a)
                self.assertIn(model_id, catalog.json()["allowed_models"])
                selected = client.post(
                    "/api/settings/model",
                    headers=user_a,
                    json={"model": model_id},
                )
                self.assertEqual(selected.status_code, 200, selected.text)
                self.assertEqual(selected.json()["model"], model_id)
                self.assertEqual(selected.json()["credential_scope"], "server_managed")
                self.assertNotEqual(client.get("/api/settings/model", headers=user_b).json()["model"], model_id)

                captured = {}

                async def fake_model_stream(client_session, endpoint, headers, request_payload, progress, stream_state):
                    captured.update({"endpoint": endpoint, "headers": headers, "payload": request_payload})
                    return ({"role": "assistant", "content": "模型作用域验证完成"}, {"total_tokens": 3}, "response-model-scope")

                with patch("app.agent._stream_completion", new=fake_model_stream):
                    chat_response = client.post("/api/chat", headers=user_a, json={"message": "查看工厂总览"})
                self.assertEqual(chat_response.status_code, 200, chat_response.text)
                self.assertEqual(captured["payload"]["model"], "enterprise-test-v1")
                self.assertEqual(captured["endpoint"], "https://models.example.test/v1/chat/completions")
                self.assertEqual(captured["headers"]["Authorization"], "Bearer server-managed-test-key")
                client.delete(f"/api/sessions/{chat_response.json()['session_id']}/permanent", headers=user_a)

                client.post("/api/settings/model", headers=user_a, json={"model": "deepseekv4-flash"})
                retired = client.patch(
                    f"/api/admin/models/{model_id}",
                    headers=admin_headers,
                    json={"status": "retired"},
                )
                self.assertEqual(retired.status_code, 200, retired.text)
                self.assertNotIn(model_id, client.get("/api/settings/model", headers=user_a).json()["allowed_models"])
        finally:
            with db_session() as connection:
                connection.execute("DELETE FROM agent.user_model_preferences WHERE model_id = ?", (model_id,))
                connection.execute("DELETE FROM agent.model_catalog WHERE model_id = ?", (model_id,))
                connection.commit()

    def test_unknown_model_is_rejected(self):
        with TestClient(app) as client:
            response = client.post("/api/settings/model", json={"model": "unknown-model"})
            self.assertEqual(response.status_code, 400)

    def test_full_logging_still_redacts_credentials(self):
        payload = sanitize({
            "DEEPSEEK_API_KEY": "secret-key",
            "headers": {"Authorization": "Bearer secret-token"},
            "input_tokens": 42,
        })
        self.assertEqual(payload["DEEPSEEK_API_KEY"], "<redacted>")
        self.assertEqual(payload["headers"]["Authorization"], "<redacted>")
        self.assertEqual(payload["input_tokens"], 42)

    def test_session_messages_survive_reload(self):
        with TestClient(app) as client:
            created = client.post("/api/sessions")
            self.assertEqual(created.status_code, 200)
            session_id = created.json()["session_id"]
            append_message(session_id, "user", "分析 CMP-03")
            append_message(session_id, "assistant", "测试回答", {"tool_trace": [{"tool": "get_equipment_health", "status": "ok", "rows": 3}]})
            restored = client.get(f"/api/sessions/{session_id}")
            self.assertEqual(restored.status_code, 200)
            self.assertEqual([item["role"] for item in restored.json()["messages"]], ["user", "assistant"])
            self.assertEqual(restored.json()["messages"][1]["metadata"]["tool_trace"][0]["rows"], 3)
            client.delete(f"/api/sessions/{session_id}")

    def test_session_history_can_list_rename_and_archive(self):
        with TestClient(app) as client:
            session = client.post("/api/sessions").json()
            append_message(session["session_id"], "user", "检查 ETCH-04 endpoint")
            listed = client.get("/api/sessions?q=ETCH-04").json()
            self.assertIn(session["session_id"], {item["session_id"] for item in listed})
            renamed = client.patch(f"/api/sessions/{session['session_id']}", json={"title": "ETCH endpoint 历史"})
            self.assertEqual(renamed.status_code, 200)
            self.assertEqual(renamed.json()["title"], "ETCH endpoint 历史")
            archived = client.delete(f"/api/sessions/{session['session_id']}")
            self.assertEqual(archived.status_code, 200)
            active_ids = {item["session_id"] for item in client.get("/api/sessions").json()}
            self.assertNotIn(session["session_id"], active_ids)
            archived_sessions = client.get("/api/sessions?include_archived=true").json()
            archived_item = next(item for item in archived_sessions if item["session_id"] == session["session_id"])
            self.assertIsNotNone(archived_item["archived_at"])
            restored = client.post(f"/api/sessions/{session['session_id']}/restore")
            self.assertEqual(restored.status_code, 200)
            self.assertIsNone(restored.json()["archived_at"])
            active_ids = {item["session_id"] for item in client.get("/api/sessions").json()}
            self.assertIn(session["session_id"], active_ids)
            active_delete = client.delete(f"/api/sessions/{session['session_id']}/permanent")
            self.assertEqual(active_delete.status_code, 200)
            self.assertEqual(client.get(f"/api/sessions/{session['session_id']}").status_code, 404)

            archived_for_delete = client.post("/api/sessions").json()
            client.delete(f"/api/sessions/{archived_for_delete['session_id']}")
            deleted = client.delete(f"/api/sessions/{archived_for_delete['session_id']}/permanent")
            self.assertEqual(deleted.status_code, 200)
            self.assertTrue(deleted.json()["deleted"])
            self.assertEqual(client.get(f"/api/sessions/{archived_for_delete['session_id']}").status_code, 404)

    def test_sessions_can_be_deleted_in_bulk_with_agent_records(self):
        with TestClient(app) as client:
            active = client.post("/api/sessions").json()
            archived = client.post("/api/sessions").json()
            untouched = client.post("/api/sessions").json()
            assistant = append_message(active["session_id"], "assistant", "批量删除级联测试")
            record_agent_result(
                active["session_id"],
                assistant["message_id"],
                {
                    "run_id": "run-bulk-delete-test",
                    "model": "test-model",
                    "mode": "test",
                    "tool_trace": [{"tool": "get_fab_overview", "status": "ok", "rows": 1}],
                },
            )
            client.delete(f"/api/sessions/{archived['session_id']}")

            missing_id = "missing-session"
            response = client.post(
                "/api/sessions/bulk-delete",
                json={"session_ids": [active["session_id"], archived["session_id"], active["session_id"], missing_id]},
            )
            self.assertEqual(response.status_code, 200)
            payload = response.json()
            self.assertEqual(payload["deleted_count"], 2)
            self.assertEqual(payload["deleted_session_ids"], [active["session_id"], archived["session_id"]])
            self.assertEqual(payload["missing_session_ids"], [missing_id])
            self.assertEqual(client.get(f"/api/sessions/{active['session_id']}").status_code, 404)
            self.assertEqual(client.get(f"/api/sessions/{archived['session_id']}").status_code, 404)
            self.assertEqual(client.get(f"/api/sessions/{untouched['session_id']}").status_code, 200)
            with db_session() as connection:
                message_count = connection.execute(
                    "SELECT COUNT(*) FROM chat_messages WHERE session_id IN (?, ?)",
                    (active["session_id"], archived["session_id"]),
                ).fetchone()[0]
                run_count = connection.execute(
                    "SELECT COUNT(*) FROM agent_runs WHERE session_id IN (?, ?)",
                    (active["session_id"], archived["session_id"]),
                ).fetchone()[0]
                tool_count = connection.execute(
                    "SELECT COUNT(*) FROM agent_tool_events WHERE run_id = ?",
                    ("run-bulk-delete-test",),
                ).fetchone()[0]
            self.assertEqual((message_count, run_count, tool_count), (0, 0, 0))
            self.assertEqual(client.post("/api/sessions/bulk-delete", json={"session_ids": []}).status_code, 422)
            client.delete(f"/api/sessions/{untouched['session_id']}/permanent")

    def test_streaming_chat_emits_progress_before_result_and_persists(self):
        plan = {
            "plan_id": "plan-stream-test",
            "schema_version": "1.0",
            "planner_version": "1.0.0",
            "status": "ready",
            "reason": "测试计划",
            "route_intent": "capability_excursion",
            "playbook_id": "PB-CAPABILITY-001",
            "playbook_version": "1.0.0",
            "required_evidence_ids": ["CAP-01"],
            "budget": {"planned_tool_calls": 1, "max_tool_calls": 2},
            "missing_inputs": [],
            "clarification_questions": [],
            "steps": [{
                "step_id": "S01",
                "title": "制程能力",
                "tool": "get_process_capability",
                "parameters": {"process_step": "CMP"},
                "parameter_sources": {"process_step": "intent_route.process_steps[0]"},
                "evidence_ids": ["CAP-01"],
                "depends_on": [],
                "required": True,
                "success_criteria": ["返回设备能力指标"],
                "estimated_result_limit": 100,
            }],
        }

        async def fake_agent(question, history, progress=None):
            await progress({"event": "run_started", "run_id": "run-test", "message": "分析任务已创建"})
            await progress({"event": "stage", "stage": "plan", "message": "执行计划已生成", "plan": plan})
            await progress({"event": "tool_started", "tool": "get_process_capability", "arguments": {"process_step": "CMP"}})
            await progress({"event": "tool_finished", "tool": "get_process_capability", "arguments": {"process_step": "CMP"}, "status": "ok", "rows": 3})
            await progress({"event": "stage", "stage": "synthesis", "message": "正在汇总结论"})
            await progress({"event": "answer_delta", "delta": "CMP "})
            await progress({"event": "answer_delta", "delta": "分析完成"})
            return {
                "answer": "CMP 分析完成",
                "tool_trace": [{"tool": "get_process_capability", "arguments": {"process_step": "CMP"}, "status": "ok", "rows": 3}],
                "sources": [],
                "mode": "test",
                "run_id": "run-test",
            }

        with patch("app.main.run_agent", new=fake_agent), TestClient(app) as client:
            created = client.post("/api/sessions").json()
            response = client.post("/api/chat/stream", json={"message": "分析 CMP", "session_id": created["session_id"]})
            self.assertEqual(response.status_code, 200)
            events = [json.loads(line) for line in response.text.splitlines() if line]
            event_types = [event["type"] for event in events]
            self.assertEqual(event_types[0], "session")
            self.assertGreaterEqual(event_types.count("progress"), 5)
            self.assertEqual(event_types[-1], "result")
            self.assertLess(event_types.index("progress"), event_types.index("result"))
            answer_events = [event for event in events if event.get("event") == "answer_delta"]
            self.assertEqual("".join(event["delta"] for event in answer_events), "CMP 分析完成")
            restored = client.get(f"/api/sessions/{created['session_id']}").json()
            self.assertEqual([item["role"] for item in restored["messages"]], ["user", "assistant"])
            self.assertEqual(restored["messages"][1]["metadata"]["run_id"], "run-test")
            execution = restored["messages"][1]["metadata"]["execution"]
            self.assertEqual(execution["execution_status"], "completed")
            self.assertEqual(execution["steps"][0]["status"], "succeeded")
            progress_tools = [item for item in events if item.get("event") in {"tool_started", "tool_finished"}]
            self.assertEqual({item.get("step_id") for item in progress_tools}, {"S01"})
            audit = client.get("/api/agent-runs/run-test/audit")
            self.assertEqual(audit.status_code, 200, audit.text)
            self.assertEqual(audit.json()["tool_events"][0]["plan_step_id"], "S01")
            self.assertTrue(audit.json()["tool_events"][0]["started_at"])
            timeline = audit.json()["timeline"]
            self.assertEqual([item["replay_sequence"] for item in timeline], list(range(1, len(timeline) + 1)))
            self.assertEqual(timeline[0]["event"], "run_started")
            self.assertEqual(timeline[-1]["event"], "run_finished")
            self.assertTrue({"plan_created", "step_started", "tool_finished", "step_finished"}.issubset(
                {item["event"] for item in timeline}
            ))
            denied = client.get(
                "/api/agent-runs/run-test/audit",
                headers={"X-Tenant-ID": "demo-fab", "X-User-ID": "other-user"},
            )
            self.assertEqual(denied.status_code, 404)
            client.delete(f"/api/sessions/{created['session_id']}")

    def test_interrupted_plan_is_recovered_with_resume_boundary(self):
        with TestClient(app) as client:
            session = client.post("/api/sessions").json()
            user_message = append_message(session["session_id"], "user", "分析 CMP 中断恢复")
            run_id = "run-interruption-recovery-test"
            start_run_record(run_id, session["session_id"], user_message["message_id"], "分析 CMP 中断恢复", "test-model")
            persist_execution_plan(run_id, {
                "plan_id": "plan-interruption-recovery-test",
                "schema_version": "1.0",
                "planner_version": "1.0.0",
                "status": "ready",
                "route_intent": "capability_excursion",
                "required_evidence_ids": ["CAP-01", "CAP-02"],
                "steps": [
                    {
                        "step_id": "S01", "title": "能力", "tool": "get_process_capability",
                        "parameters": {"process_step": "CMP"}, "evidence_ids": ["CAP-01"],
                        "required": True,
                    },
                    {
                        "step_id": "S02", "title": "异常", "tool": "get_recent_excursions",
                        "parameters": {"process_step": "CMP"}, "evidence_ids": ["CAP-02"],
                        "depends_on": ["S01"], "required": True,
                    },
                ],
            })
            start_tool_event(run_id, "get_process_capability", {"process_step": "CMP"})

            recovered = recover_interrupted_executions("测试进程重启")
            self.assertGreaterEqual(recovered, 1)
            audit = client.get(f"/api/agent-runs/{run_id}/audit")
            self.assertEqual(audit.status_code, 200, audit.text)
            payload = audit.json()
            self.assertEqual(payload["run"]["status"], "interrupted")
            self.assertEqual(payload["execution"]["execution_status"], "interrupted")
            self.assertEqual(payload["execution"]["steps"][0]["status"], "interrupted")
            self.assertEqual(payload["execution"]["steps"][1]["status"], "pending")
            self.assertTrue(payload["execution"]["recovery"]["resumable"])
            self.assertEqual(payload["execution"]["recovery"]["resume_from_step_ids"], ["S01", "S02"])
            self.assertEqual(payload["tool_events"][0]["status"], "interrupted")
            client.delete(f"/api/sessions/{session['session_id']}/permanent")

    def test_admin_overview_tracks_user_agent_runtime(self):
        async def fake_agent(question, history, progress=None):
            await progress({"event": "run_started", "run_id": "run-admin-runtime", "message": "分析任务已创建"})
            await progress({"event": "answer_delta", "delta": "完成"})
            return {
                "answer": "运行状态测试完成",
                "tool_trace": [],
                "sources": [],
                "mode": "test",
                "model": "runtime-test-model",
                "run_id": "run-admin-runtime",
            }

        admin_token = "runtime-admin-token"
        user_headers = {"X-Tenant-ID": "runtime-fab", "X-User-ID": "operator-01"}
        with (
            patch.dict(os.environ, {"FAB_ADMIN_API_TOKEN": admin_token}),
            patch("app.main.run_agent", new=fake_agent),
            TestClient(app) as client,
        ):
            response = client.post("/api/chat/stream", headers=user_headers, json={"message": "检查运行状态"})
            self.assertEqual(response.status_code, 200, response.text)
            events = [json.loads(line) for line in response.text.splitlines() if line]
            session_id = next(item["session_id"] for item in events if item["type"] == "session")
            renamed = client.patch(
                f"/api/sessions/{session_id}",
                headers=user_headers,
                json={"title": "用户自定义的设备调查会话"},
            )
            self.assertEqual(renamed.status_code, 200, renamed.text)
            overview = client.get(
                "/api/admin/overview?hours=24",
                headers={"Authorization": f"Bearer {admin_token}"},
            )
            self.assertEqual(overview.status_code, 200, overview.text)
            user = next(
                item for item in overview.json()["users"]
                if item["tenant_id"] == "runtime-fab" and item["user_id"] == "operator-01"
            )
            self.assertEqual(user["status"], "idle")
            self.assertEqual(user["active_runs"], 0)
            self.assertEqual(user["model"], "runtime-test-model")
            self.assertGreaterEqual(user["run_count"], 1)
            recent_run = next(
                item for item in overview.json()["recent_runs"]
                if item["run_id"] == "run-admin-runtime"
            )
            self.assertEqual(recent_run["question_preview"], "检查运行状态")
            self.assertEqual(recent_run["session_id"], session_id)
            self.assertEqual(recent_run["session_title"], "用户自定义的设备调查会话")
            self.assertIsInstance(recent_run["user_message_id"], int)
            with db_session() as connection:
                persisted = connection.execute(
                    """SELECT r.question_preview, m.content, m.role
                       FROM agent_runs r JOIN chat_messages m ON m.message_id = r.user_message_id
                       WHERE r.run_id = ?""",
                    ("run-admin-runtime",),
                ).fetchone()
            self.assertEqual(tuple(persisted), ("检查运行状态", "检查运行状态", "user"))
            client.delete(f"/api/sessions/{session_id}/permanent", headers=user_headers)

    def test_stream_retry_with_same_client_message_id_is_idempotent(self):
        calls = 0

        async def fake_agent(question, history, progress=None):
            nonlocal calls
            calls += 1
            await progress({"event": "run_started", "run_id": "run-idempotent", "message": "分析任务已创建"})
            await progress({"event": "answer_delta", "delta": "完成"})
            return {
                "answer": "完成",
                "tool_trace": [],
                "sources": [],
                "mode": "test",
                "model": "test-model",
                "run_id": "run-idempotent",
            }

        payload_id = "msg-idempotent-0001"
        with patch("app.main.run_agent", new=fake_agent), TestClient(app) as client:
            session_id = client.post("/api/sessions").json()["session_id"]
            payload = {"message": "分析 CMP", "session_id": session_id, "client_message_id": payload_id}
            first = client.post("/api/chat/stream", json=payload)
            second = client.post("/api/chat/stream", json=payload)
            self.assertEqual(first.status_code, 200)
            self.assertEqual(second.status_code, 200)
            second_events = [json.loads(line) for line in second.text.splitlines() if line]
            self.assertTrue(second_events[-1]["result"]["cached"])
            self.assertEqual(calls, 1)
            conflict = client.post("/api/chat/stream", json={**payload, "message": "分析 ETCH"})
            self.assertEqual(conflict.status_code, 409)
            self.assertEqual(conflict.json()["error_code"], "IDEMPOTENCY_CONFLICT")
            restored = client.get(f"/api/sessions/{session_id}").json()
            self.assertEqual([item["role"] for item in restored["messages"]], ["user", "assistant"])
            client.delete(f"/api/sessions/{session_id}/permanent")

    def test_model_retries_transient_500_before_first_answer_token(self):
        attempts = 0
        progress_events = []
        request = httpx.Request("POST", "https://example.test/chat/completions")
        transient_error = httpx.HTTPStatusError(
            "temporary failure",
            request=request,
            response=httpx.Response(500, request=request),
        )

        async def fake_stream(*args):
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                raise transient_error
            return ({"role": "assistant", "content": "OK"}, {"total_tokens": 3}, "response-test")

        async def progress(event):
            progress_events.append(event)

        settings = SimpleNamespace(
            deepseek_api_key="test-key",
            deepseek_base_url="https://example.test",
            deepseek_model="test-model",
            api_timeout_seconds=5,
            model_max_retries=2,
            model_retry_base_seconds=0.05,
        )
        with (
            patch("app.agent.get_settings", return_value=settings),
            patch("app.agent._stream_completion", new=fake_stream),
            patch("app.agent.asyncio.sleep", new=AsyncMock()),
        ):
            result = asyncio.run(agent_module._run_agent_impl("查看工厂总览", [], "run-retry-test", progress))
        self.assertEqual(result["answer"], "OK")
        self.assertEqual(attempts, 2)
        self.assertTrue(any(item.get("stage") == "retrying" for item in progress_events))

    def test_model_does_not_retry_after_answer_output_starts(self):
        attempts = 0
        request = httpx.Request("POST", "https://example.test/chat/completions")
        stream_error = httpx.HTTPStatusError(
            "stream interrupted",
            request=request,
            response=httpx.Response(500, request=request),
        )

        async def failing_stream(*args):
            nonlocal attempts
            attempts += 1
            args[-1]["answer_emitted"] = True
            raise stream_error

        settings = SimpleNamespace(
            deepseek_api_key="test-key",
            deepseek_base_url="https://example.test",
            deepseek_model="test-model",
            api_timeout_seconds=5,
            model_max_retries=2,
            model_retry_base_seconds=0.05,
        )
        with (
            patch("app.agent.get_settings", return_value=settings),
            patch("app.agent._stream_completion", new=failing_stream),
            patch("app.agent.asyncio.sleep", new=AsyncMock()),
        ):
            with self.assertRaises(httpx.HTTPStatusError):
                asyncio.run(agent_module._run_agent_impl("查看工厂总览", [], "run-no-retry-test"))
        self.assertEqual(attempts, 1)

    def test_stream_error_is_structured_and_audited(self):
        request_id = "req-stream-error-001"
        upstream_request = httpx.Request("POST", "https://example.test/chat/completions")

        async def failing_agent(*args, **kwargs):
            raise httpx.HTTPStatusError(
                "upstream unavailable",
                request=upstream_request,
                response=httpx.Response(503, request=upstream_request),
            )

        with patch("app.main.run_agent", new=failing_agent), TestClient(app) as client:
            session_id = client.post("/api/sessions").json()["session_id"]
            response = client.post(
                "/api/chat/stream",
                json={"message": "分析 CMP", "session_id": session_id, "client_message_id": "msg-stream-error-001"},
                headers={"X-Request-ID": request_id},
            )
            events = [json.loads(line) for line in response.text.splitlines() if line]
            self.assertEqual(response.status_code, 200)
            self.assertEqual(events[-1]["type"], "error")
            self.assertEqual(events[-1]["error"]["error_code"], "MODEL_UPSTREAM_ERROR")
            self.assertEqual(events[-1]["error"]["request_id"], request_id)
            self.assertTrue(events[-1]["error"]["retryable"])
            self.assertTrue(any(
                event.get("event") == "api_stream_error" and event.get("request_id") == request_id
                for event in read_events(200)
            ))
            client.delete(f"/api/sessions/{session_id}/permanent")


if __name__ == "__main__":
    unittest.main()
