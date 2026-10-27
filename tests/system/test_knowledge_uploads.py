from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from tests.system.support import PROJECT_DIR, isolated_system_environment

isolated_system_environment()

from fastapi.testclient import TestClient

from app.agent_runtime import service as agent_module
from app.persistence.database import db_session, get_active_defect_aliases, initialize, replace_defect_type_aliases
from app.main import app
from app.knowledge.rag import _infer_filters, search_knowledge
from app.manufacturing.reports import safe_select
from app.manufacturing.seed import seed_reference_data

class KnowledgeUploadsSystemTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        seed_reference_data()

    def test_rag_returns_cited_source(self):
        results = search_knowledge("CMP 划伤 Pad Conditioner", 10)
        self.assertTrue(results)
        self.assertEqual({item["source_type"] for item in results}, {"knowledge", "case"})
        self.assertTrue(all(item["reference_id"] for item in results))
        self.assertTrue(all(item["retrieval"]["selection_reason"] for item in results))

    def test_rag_infers_chinese_process_and_defect_aliases(self):
        results = search_knowledge("刻蚀设备 CD 偏移的潜在原因", 6)
        self.assertTrue(results)
        self.assertEqual({item["source_type"] for item in results}, {"knowledge", "case"})
        self.assertTrue(all(item["retrieval"]["filters"]["process_step"] == "ETCH" for item in results))
        self.assertTrue(all(item["retrieval"]["filters"]["defect_type"] == "CD Shift" for item in results))

    def test_rag_uses_database_managed_defect_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "aliases.db"
            initialize(database_path)
            replace_defect_type_aliases("Pattern Collapse", ("pattern collapse", "图形倒塌"), database_path)
            aliases = get_active_defect_aliases(database_path)
        with patch("app.knowledge.rag.get_active_defect_aliases", return_value=aliases):
            filters = _infer_filters("光刻后出现图形倒塌", None, None, None)
        self.assertEqual(filters["defect_type"], "Pattern Collapse")

    def test_rag_gold_dataset(self):
        dataset_path = PROJECT_DIR / "evaluations" / "rag_gold.json"
        dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
        for item in dataset:
            with self.subTest(item=item["id"]):
                results = search_knowledge(item["query"], limit=6, **item.get("filters", {}))
                references = {result["reference_id"] for result in results}
                source_types = {result["source_type"] for result in results}
                self.assertTrue(references.intersection(item["expected_any"]))
                self.assertTrue(set(item.get("required_source_types", [])).issubset(source_types))
                self.assertFalse(references.intersection(item.get("forbidden", [])))

    def test_knowledge_documents_support_chunk_navigation(self):
        with TestClient(app) as client:
            documents = client.get("/api/knowledge/documents")
            self.assertEqual(documents.status_code, 200)
            self.assertGreaterEqual(len(documents.json()), 12)
            first_chunk_id = documents.json()[0]["first_chunk_id"]
            document = client.get(f"/api/knowledge/chunks/{first_chunk_id}")
            self.assertEqual(document.status_code, 200)
            self.assertEqual(document.json()["selected_chunk_id"], first_chunk_id)
            self.assertTrue(document.json()["chunks"])
            self.assertTrue(all("content" in chunk for chunk in document.json()["chunks"]))
            self.assertEqual(document.json()["document"]["status"], "Active")
            self.assertTrue(document.json()["document"]["version"])
            self.assertGreaterEqual(document.json()["document"]["authority"], 1)
            self.assertEqual(client.get("/api/knowledge/chunks/999999999").status_code, 404)

    def test_historical_cases_keep_validated_catalog_and_missing_case_404(self):
        with TestClient(app) as client:
            cases = client.get("/api/cases")
            self.assertEqual(cases.status_code, 200)
            self.assertTrue(cases.json())
            self.assertTrue(all(case["case_status"] == "Validated" for case in cases.json()))
            first_case_id = cases.json()[0]["case_id"]
            detail = client.get(f"/api/cases/{first_case_id}")
            self.assertEqual(detail.status_code, 200)
            self.assertEqual(detail.json()["case_id"], first_case_id)
            self.assertEqual(detail.json()["case_status"], "Validated")
            self.assertEqual(client.get("/api/cases/CASE-DOES-NOT-EXIST").status_code, 404)

    def test_personal_knowledge_upload_search_privacy_and_delete(self):
        content = (
            "# CMP-03 Pad 更换记录\n\n"
            "2026-07-01 完成 Pad 更换。更换后确认主轴电流稳定，抛光压力恢复正常。\n\n"
            "## 验证结果\n\n"
            "连续三个批次膜厚均匀性满足要求，Scratch 缺陷未再次出现。"
        ).encode("utf-8")
        headers = {"Idempotency-Key": "personal-upload-test-001"}
        metadata = json.dumps({
            "title": "CMP-03 Pad 更换记录",
            "process_steps": ["CMP"],
            "equipment_ids": ["CMP-03"],
            "effective_date": "2026-07-01",
            "tags": ["pad", "maintenance"],
        }, ensure_ascii=False)

        with TestClient(app) as client:
            uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers=headers,
                files={"file": ("cmp-pad.md", content, "text/markdown")},
                data={"metadata": metadata},
            )
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            document_id = uploaded.json()["document_id"]
            job_id = uploaded.json()["job_id"]

            job = client.get(f"/api/knowledge/jobs/{job_id}")
            self.assertEqual(job.status_code, 200)
            self.assertEqual(job.json()["status"], "ready")
            self.assertEqual(job.json()["progress"], 100)

            listed = client.get("/api/knowledge/personal/documents?status=ready&q=CMP")
            self.assertEqual(listed.status_code, 200)
            self.assertIn(document_id, {item["document_id"] for item in listed.json()["items"]})

            detail = client.get(f"/api/knowledge/personal/documents/{document_id}")
            self.assertEqual(detail.status_code, 200)
            self.assertEqual(detail.json()["process_steps"], ["CMP"])
            self.assertEqual(detail.json()["tags_source"], "user")
            self.assertGreater(detail.json()["chunk_count"], 0)
            self.assertNotIn("owner_user_id", detail.json())
            self.assertNotIn("storage_key", detail.json())

            document_content = client.get(f"/api/knowledge/personal/documents/{document_id}/content")
            self.assertEqual(document_content.status_code, 200)
            self.assertTrue(document_content.json()["chunks"])
            chunk_id = document_content.json()["chunks"][0]["chunk_id"]
            with db_session() as connection:
                indexed_chunks = connection.execute(
                    "SELECT COUNT(*) FROM personal_knowledge_fts WHERE document_id = ?", (document_id,)
                ).fetchone()[0]
            self.assertEqual(indexed_chunks, detail.json()["chunk_count"])

            search = client.post(
                "/api/knowledge/personal/search-preview",
                json={"query": "CMP-03 Pad 更换后 Scratch 验证", "document_ids": [document_id]},
            )
            self.assertEqual(search.status_code, 200, search.text)
            self.assertGreater(search.json()["count"], 0)
            self.assertTrue(all(item["source_type"] == "personal" for item in search.json()["results"]))

            chunk = client.get(f"/api/knowledge/personal/chunks/{chunk_id}")
            self.assertEqual(chunk.status_code, 200)
            self.assertEqual(chunk.json()["document_id"], document_id)

            updated = client.patch(
                f"/api/knowledge/personal/documents/{document_id}",
                json={"title": "CMP-03 Pad 更换与验证记录", "tags": ["pad", "verified"]},
            )
            self.assertEqual(updated.status_code, 200, updated.text)
            self.assertEqual(updated.json()["title"], "CMP-03 Pad 更换与验证记录")
            self.assertEqual(updated.json()["tags_source"], "user")

            tag_only = client.post(
                "/api/knowledge/personal/search-preview",
                json={"query": "verified", "document_ids": [document_id]},
            )
            self.assertEqual(tag_only.status_code, 200, tag_only.text)
            self.assertIn(document_id, {item["document_id"] for item in tag_only.json()["results"]})

            # Exercise the retry API after a simulated retryable failure: user tags
            # must survive a fresh extraction and index rebuild.
            with db_session() as connection:
                connection.execute(
                    """UPDATE knowledge.personal_knowledge_documents
                       SET status = 'failed', error_code = 'INGESTION_FAILED'
                       WHERE document_id = ?""",
                    (document_id,),
                )
                connection.commit()
            retried = client.post(f"/api/knowledge/personal/documents/{document_id}/retry")
            self.assertEqual(retried.status_code, 202, retried.text)
            after_retry = client.get(f"/api/knowledge/personal/documents/{document_id}")
            self.assertEqual(after_retry.json()["tags"], ["pad", "verified"])
            self.assertEqual(after_retry.json()["tags_source"], "user")

            auto_uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "personal-auto-tags-003"},
                files={"file": ("note.md", "# 生产记录\n\nETCH-09 出现划伤，完成维护。".encode("utf-8"), "text/markdown")},
                data={"metadata": json.dumps({"title": "生产记录"}, ensure_ascii=False)},
            )
            self.assertEqual(auto_uploaded.status_code, 202, auto_uploaded.text)
            auto_document_id = auto_uploaded.json()["document_id"]
            auto_detail = client.get(f"/api/knowledge/personal/documents/{auto_document_id}")
            self.assertEqual(auto_detail.json()["tags_source"], "auto")
            self.assertTrue({"ETCH", "ETCH-09", "RECORD", "SCRATCH", "MAINTENANCE", "MD"}.issubset(auto_detail.json()["tags"]))
            auto_tag_only = client.post(
                "/api/knowledge/personal/search-preview",
                json={"query": "MAINTENANCE", "document_ids": [auto_document_id]},
            )
            self.assertEqual(auto_tag_only.status_code, 200, auto_tag_only.text)
            self.assertEqual([item["document_id"] for item in auto_tag_only.json()["results"]], [auto_document_id])

            empty_uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "personal-empty-tags-004"},
                files={"file": ("calibration.md", b"# Calibration\n\nBaseline complete.", "text/markdown")},
                data={"metadata": '{"tags": []}'},
            )
            self.assertEqual(empty_uploaded.status_code, 202, empty_uploaded.text)
            empty_document_id = empty_uploaded.json()["document_id"]
            empty_detail = client.get(f"/api/knowledge/personal/documents/{empty_document_id}")
            self.assertEqual(empty_detail.json()["tags_source"], "user")
            self.assertEqual(empty_detail.json()["tags"], [])
            sketch = client.patch(
                f"/api/knowledge/personal/documents/{empty_document_id}", json={"tags": ["sketch"]}
            )
            self.assertEqual(sketch.status_code, 200, sketch.text)
            self.assertEqual(client.post(
                "/api/knowledge/personal/search-preview",
                json={"query": "etch", "document_ids": [empty_document_id]},
            ).json()["count"], 0)

            downloaded = client.get(f"/api/knowledge/personal/documents/{document_id}/file")
            self.assertEqual(downloaded.status_code, 200)
            self.assertEqual(downloaded.content, content)

            replay = client.post(
                "/api/knowledge/personal/documents",
                headers=headers,
                files={"file": ("cmp-pad.md", content, "text/markdown")},
                data={"metadata": metadata},
            )
            self.assertEqual(replay.status_code, 202)
            self.assertEqual(replay.json()["document_id"], document_id)
            self.assertEqual(replay.json()["job_id"], job_id)

            duplicate = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "personal-upload-test-002"},
                files={"file": ("duplicate.md", content, "text/markdown")},
            )
            self.assertEqual(duplicate.status_code, 409)
            self.assertEqual(duplicate.json()["error_code"], "DUPLICATE_DOCUMENT")

            with patch.dict(os.environ, {"FAB_PERSONAL_KNOWLEDGE_USER_ID": "different-user"}):
                self.assertEqual(client.get(f"/api/knowledge/personal/documents/{document_id}").status_code, 404)
                self.assertEqual(client.get(f"/api/knowledge/personal/chunks/{chunk_id}").status_code, 404)
                self.assertEqual(client.get(f"/api/knowledge/jobs/{job_id}").status_code, 404)
                private_search = client.post(
                    "/api/knowledge/personal/search-preview",
                    json={"query": "CMP Pad", "document_ids": [document_id]},
                )
                self.assertEqual(private_search.status_code, 404)

            deleting = client.delete(f"/api/knowledge/personal/documents/{document_id}")
            self.assertEqual(deleting.status_code, 202, deleting.text)
            delete_job_id = deleting.json()["job_id"]
            self.assertEqual(client.get(f"/api/knowledge/jobs/{delete_job_id}").json()["status"], "deleted")
            self.assertEqual(client.get(f"/api/knowledge/personal/documents/{document_id}").status_code, 404)
            self.assertEqual(client.get(f"/api/knowledge/personal/chunks/{chunk_id}").status_code, 404)
            with db_session() as connection:
                indexed_chunks = connection.execute(
                    "SELECT COUNT(*) FROM personal_knowledge_fts WHERE document_id = ?", (document_id,)
                ).fetchone()[0]
            self.assertEqual(indexed_chunks, 0)
            for extra_document_id in (auto_document_id, empty_document_id):
                deleted_extra = client.delete(f"/api/knowledge/personal/documents/{extra_document_id}")
                self.assertEqual(deleted_extra.status_code, 202, deleted_extra.text)
                self.assertEqual(client.get(f"/api/knowledge/personal/documents/{extra_document_id}").status_code, 404)

    def test_agent_search_combines_private_knowledge_with_cross_language_recall(self):
        content = (
            "# Introduction to Photolithography\n\n"
            "Photolithography transfers a reticle pattern into photoresist by optical exposure.\n\n"
            "## Process basics\n\n"
            "The basic sequence covers coat, soft bake, align, expose, develop, inspect, and pattern transfer."
        ).encode("utf-8")
        metadata = json.dumps({
            "title": "Introduction to Photolithography",
            "process_steps": ["LITHO"],
            "tags": ["training", "fundamentals"],
        })
        with TestClient(app) as client:
            uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "agent-personal-rag-test-001"},
                files={"file": ("intro-photolithography.md", content, "text/markdown")},
                data={"metadata": metadata},
            )
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            document_id = uploaded.json()["document_id"]
            try:
                results = agent_module.search_agent_knowledge(
                    "请介绍光刻技术原理",
                    limit=6,
                    source_types=["personal"],
                )
                self.assertTrue(results)
                self.assertTrue(all(item["source_type"] == "personal" for item in results))
                self.assertEqual(results[0]["document_id"], document_id)
                self.assertEqual(results[0]["source"], "个人知识库")
                self.assertEqual(results[0]["process_steps"], ["LITHO"])
                self.assertTrue(results[0]["reference_id"].startswith("PERSONAL-"))
                renamed = client.patch(
                    f"/api/knowledge/personal/documents/{document_id}",
                    json={"title": "我的光刻培训资料"},
                )
                self.assertEqual(renamed.status_code, 200, renamed.text)
                self.assertEqual(renamed.json()["title"], "我的光刻培训资料")
                renamed_results = agent_module.search_agent_knowledge(
                    "请介绍光刻技术原理",
                    limit=6,
                    source_types=["personal"],
                )
                self.assertEqual(renamed_results[0]["title"], "我的光刻培训资料")
                self.assertNotIn(
                    "personal",
                    {item["source_type"] for item in search_knowledge("请介绍光刻技术原理", limit=6)},
                )
            finally:
                deleted = client.delete(f"/api/knowledge/personal/documents/{document_id}")
                self.assertEqual(deleted.status_code, 202, deleted.text)

    def test_personal_knowledge_upload_validation_and_token(self):
        with TestClient(app) as client:
            missing_key = client.post(
                "/api/knowledge/personal/documents",
                files={"file": ("note.txt", b"valid text content", "text/plain")},
            )
            self.assertEqual(missing_key.status_code, 422)

            unsupported = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "unsupported-file-test"},
                files={"file": ("payload.exe", b"not executable", "application/octet-stream")},
            )
            self.assertEqual(unsupported.status_code, 415)
            self.assertEqual(unsupported.json()["error_code"], "UNSUPPORTED_FILE_TYPE")

            with patch.dict(os.environ, {"FAB_PERSONAL_KNOWLEDGE_API_TOKEN": "private-test-token"}):
                unauthenticated = client.get("/api/knowledge/personal/documents")
                self.assertEqual(unauthenticated.status_code, 401)
                authenticated = client.get(
                    "/api/knowledge/personal/documents",
                    headers={"Authorization": "Bearer private-test-token"},
                )
                self.assertEqual(authenticated.status_code, 200)

    def test_personal_knowledge_blank_scanned_pdf_fails_without_silent_indexing(self):
        from pypdf import PdfWriter

        output = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(output)
        with TestClient(app) as client:
            uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "blank-ocr-test"},
                files={"file": ("scanned.pdf", output.getvalue(), "application/pdf")},
            )
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            job = client.get(f"/api/knowledge/jobs/{uploaded.json()['job_id']}")
            self.assertEqual(job.status_code, 200)
            self.assertEqual(job.json()["status"], "failed")
            self.assertEqual(job.json()["error"]["error_code"], "INVALID_FILE_CONTENT")
            deleted = client.delete(f"/api/knowledge/personal/documents/{uploaded.json()['document_id']}")
            self.assertEqual(deleted.status_code, 202)

    def test_personal_knowledge_low_quality_image_is_not_indexed(self):
        from PIL import Image
        from app.knowledge.ocr import OCRResult

        buffer = io.BytesIO()
        Image.new("RGB", (80, 40), "white").save(buffer, "PNG")
        low = OCRResult("barely", (0.1,), 1, 0.1, 1.0, True)
        with patch(
            "app.personal_knowledge_runtime.service.extract_image_result",
            return_value=low,
        ), TestClient(app) as client:
            uploaded = client.post(
                "/api/knowledge/personal/documents",
                headers={"Idempotency-Key": "low-quality-ocr-test"},
                files={"file": ("low.png", buffer.getvalue(), "image/png")},
            )
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            document_id, job_id = uploaded.json()["document_id"], uploaded.json()["job_id"]
            try:
                job = client.get(f"/api/knowledge/jobs/{job_id}")
                self.assertEqual(job.json()["status"], "failed")
                self.assertEqual(job.json()["error"]["error_code"], "OCR_QUALITY_LOW")
                with db_session() as connection:
                    chunk_count = connection.execute(
                        "SELECT COUNT(*) FROM personal_knowledge_chunks "
                        "WHERE document_id = ?",
                        (document_id,),
                    ).fetchone()[0]
                    fts_count = connection.execute(
                        "SELECT COUNT(*) FROM personal_knowledge_fts "
                        "WHERE document_id = ?",
                        (document_id,),
                    ).fetchone()[0]
                    self.assertEqual(chunk_count, 0)
                    self.assertEqual(fts_count, 0)
                preview = client.post(
                    "/api/knowledge/personal/search-preview",
                    json={"query": "barely", "limit": 10},
                )
                self.assertEqual(preview.status_code, 200, preview.text)
                self.assertNotIn(
                    document_id,
                    {item["document_id"] for item in preview.json()["results"]},
                )
            finally:
                deleted = client.delete(
                    f"/api/knowledge/personal/documents/{document_id}"
                )
                self.assertEqual(deleted.status_code, 202)

    def test_personal_knowledge_office_and_csv_parsers(self):
        from docx import Document
        from openpyxl import Workbook

        docx_buffer = io.BytesIO()
        docx_document = Document()
        docx_document.add_heading("ETCH-01 维护记录", level=1)
        docx_document.add_paragraph("腔体清洗完成，Particle 复测符合要求。")
        docx_document.save(docx_buffer)

        xlsx_buffer = io.BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CVD-02 检查"
        sheet.append(["项目", "结果"])
        sheet.append(["膜厚均匀性", "通过"])
        workbook.save(xlsx_buffer)
        workbook.close()

        samples = (
            ("etch-maintenance.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "腔体清洗"),
            ("cvd-check.xlsx", xlsx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "膜厚均匀性"),
            ("cmp-check.csv", "项目,结果\nPad检查,通过\n".encode("utf-8"), "text/csv", "Pad检查"),
        )
        with TestClient(app) as client:
            for index, (filename, content, mime_type, expected) in enumerate(samples):
                with self.subTest(filename=filename):
                    uploaded = client.post(
                        "/api/knowledge/personal/documents",
                        headers={"Idempotency-Key": f"parser-format-test-{index}"},
                        files={"file": (filename, content, mime_type)},
                    )
                    self.assertEqual(uploaded.status_code, 202, uploaded.text)
                    document_id = uploaded.json()["document_id"]
                    job = client.get(f"/api/knowledge/jobs/{uploaded.json()['job_id']}")
                    self.assertEqual(job.json()["status"], "ready", job.text)
                    parsed = client.get(f"/api/knowledge/personal/documents/{document_id}/content")
                    self.assertIn(expected, " ".join(chunk["content"] for chunk in parsed.json()["chunks"]))
                    self.assertEqual(client.delete(f"/api/knowledge/personal/documents/{document_id}").status_code, 202)

    def test_historical_case_has_fixed_format(self):
        with TestClient(app) as client:
            cases = client.get("/api/cases")
            self.assertEqual(cases.status_code, 200)
            self.assertGreaterEqual(len(cases.json()), 12)
            case_id = cases.json()[0]["case_id"]
            detail = client.get(f"/api/cases/{case_id}")
            self.assertEqual(detail.status_code, 200)
            required = {
                "problem", "solution", "result", "summary", "root_cause",
                "root_cause_confidence", "validation_method", "recurrence_status",
                "approved_by", "case_status",
            }
            self.assertTrue(required.issubset(detail.json()))
            self.assertTrue(all(detail.json()[field] for field in required))

    def test_sql_guard_rejects_writes(self):
        with self.assertRaises(ValueError):
            safe_select("DELETE FROM lots")
        with self.assertRaises(ValueError):
            safe_select("SELECT * FROM sqlite_master")
        with self.assertRaises(ValueError):
            safe_select("SELECT scenario_id FROM process_runs")

if __name__ == "__main__":
    unittest.main()
